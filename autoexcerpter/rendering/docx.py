"""DOCX summary document creation for AutoExcerpter.

Generates formatted Word documents with structured summaries, including:
- LaTeX formulas as native Word equations (see ``rendering.equations``)
- Citation deduplication with OpenAlex hyperlinks
- Document structure overview with page ranges
- XML sanitization for safe DOCX output
"""

from __future__ import annotations

import contextlib
import logging
import os
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from autoexcerpter.constants import (
    BODY_FONT_NAME,
    BODY_FONT_PT,
    BODY_SPACE_AFTER_PT,
    BULLET_FONT_PT,
    BULLET_HANGING_INDENT_CM,
    BULLET_LEFT_INDENT_CM,
    BULLET_SPACE_AFTER_PT,
    COLOR_BLACK,
    COLOR_METADATA_GRAY,
    COLOR_PAGE_HEADING,
    COLOR_REF_META_GRAY,
    COLOR_SECTION_RULE,
    FOOTER_FONT_PT,
    METADATA_FONT_PT,
    PAGE_HEADING_FONT_PT,
    PAGE_HEADING_SPACE_AFTER_PT,
    PAGE_HEADING_SPACE_BEFORE_PT,
    PAGE_HEIGHT_CM,
    PAGE_MARGIN_CM,
    PAGE_WIDTH_CM,
    REF_FONT_PT,
    REF_HANGING_INDENT_CM,
    REF_META_FONT_PT,
    REF_SPACE_AFTER_PT,
    SECTION_HEADING_FONT_PT,
    SECTION_HEADING_SPACE_AFTER_PT,
    SECTION_HEADING_SPACE_BEFORE_PT,
    TITLE_FONT_PT,
    TITLE_SPACE_AFTER_PT,
)
from autoexcerpter.rendering.citations import (
    Citation,
    CitationManager,
    enrich_if_enabled,
)
from autoexcerpter.rendering.equations import (
    add_math_to_paragraph,
    parse_latex_in_text,
)
from autoexcerpter.rendering.summary import (
    PAGE_TYPE_LABELS,
    STRUCTURE_PAGE_TYPE_ORDER,
    PageRenderData,
    SummaryData,
    collapse_internal_newlines,
    format_structure_page_range,
    prepare_summary_data,
    sanitize_for_xml,
)

logger = logging.getLogger(__name__)


def add_formatted_text_to_paragraph(paragraph: Any, text: str) -> None:
    """Add text to a paragraph, parsing and rendering LaTeX formulas."""
    segments = parse_latex_in_text(text)

    for content, segment_type in segments:
        if segment_type in ("latex_display", "latex_inline"):
            add_math_to_paragraph(
                paragraph, content, display=segment_type == "latex_display"
            )
        else:
            # Markdown **bold** and *italic* become run formatting rather than
            # literal asterisks.
            for kind, piece in parse_markdown_emphasis(content):
                run = paragraph.add_run(sanitize_for_xml(piece))
                if kind in ("bold", "bold_italic"):
                    run.bold = True
                if kind in ("italic", "bold_italic"):
                    run.italic = True


# The italic alternative requires non-space flanking characters, as CommonMark
# does: "a * b * c" is literal prose, not emphasis. Bold (**) is unchanged.
_MARKDOWN_EMPHASIS_PATTERN = re.compile(
    r"\*\*\*(?P<bold_italic>[^*]+)\*\*\*"
    r"|\*\*(?P<bold>[^*]+)\*\*"
    r"|\*(?P<italic>[^*\s](?:[^*]*[^*\s])?)\*"
)


def parse_markdown_emphasis(text: str) -> list[tuple[str, str]]:
    """Split text into ``(kind, content)`` segments by Markdown emphasis.

    *kind* is ``"text"``, ``"bold"``, ``"italic"`` or ``"bold_italic"``.
    Citation strings extracted by the LLM occasionally carry Markdown emphasis
    (e.g. ``*The American Economic Review, 87*(2)``), which the DOCX writer
    maps to run formatting.
    """
    segments: list[tuple[str, str]] = []
    last_end = 0
    for match in _MARKDOWN_EMPHASIS_PATTERN.finditer(text):
        if match.start() > last_end:
            segments.append(("text", text[last_end : match.start()]))
        if match.group("bold_italic") is not None:
            segments.append(("bold_italic", match.group("bold_italic")))
        elif match.group("bold") is not None:
            segments.append(("bold", match.group("bold")))
        else:
            segments.append(("italic", match.group("italic")))
        last_end = match.end()
    if last_end < len(text):
        segments.append(("text", text[last_end:]))
    return segments or [("text", text)]


def strip_markdown_emphasis(text: str) -> str:
    """Remove Markdown emphasis markers, keeping the emphasized text."""
    return _MARKDOWN_EMPHASIS_PATTERN.sub(
        lambda m: m.group("bold_italic") or m.group("bold") or m.group("italic"),
        text,
    )


def add_hyperlink(
    paragraph: Any, url: str, text: str, size_pt: float | None = None
) -> None:
    """Add a hyperlink to a paragraph in a DOCX document.

    When *size_pt* is set, the run is pinned to that point size (via ``w:sz``
    /``w:szCs`` half-points) so it matches sibling runs instead of inheriting
    the Normal style's default size.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "Hyperlink")
    rPr.append(rStyle)
    if size_pt is not None:
        half_points = str(int(size_pt * 2))
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), half_points)
        rPr.append(sz)
        szCs = OxmlElement("w:szCs")
        szCs.set(qn("w:val"), half_points)
        rPr.append(szCs)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)

    paragraph._p.append(hyperlink)


def _rgb(color: int) -> RGBColor:
    """Build an :class:`RGBColor` from a ``0xRRGGBB`` integer."""
    return RGBColor((color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF)


def _set_east_asian_font(element: Any, name: str) -> None:
    """Pin the East Asian font so Word does not substitute a different face.

    Also strips the theme-font attributes that the built-in Title/Heading styles
    carry on their ``w:rFonts`` element. Word and LibreOffice give ``*Theme``
    attributes precedence over an explicit ``w:ascii``/``w:hAnsi`` face, so
    without this the headings would render in the theme's major font (Calibri/
    Carlito) instead of the requested body face. ``w:cs`` is set for
    completeness so complex-script runs match too.
    """
    rpr = element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    for theme_attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{theme_attr}")
        if rfonts.get(key) is not None:
            del rfonts.attrib[key]


def _style_font(
    style: Any,
    *,
    size_pt: float,
    bold: bool | None = None,
    color: int | None = None,
    name: str = BODY_FONT_NAME,
) -> None:
    """Apply an explicit font to a named style (never rely on docx defaults)."""
    font = style.font
    font.name = name
    font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = _rgb(color)
        # Strip inherited theme colors so the explicit RGB always wins in Word.
        rpr = style.element.get_or_add_rPr()
        color_el = rpr.find(qn("w:color"))
        if color_el is not None:
            for attr in ("themeColor", "themeTint", "themeShade"):
                key = qn(f"w:{attr}")
                if color_el.get(key) is not None:
                    del color_el.attrib[key]
    _set_east_asian_font(style.element, name)


def _clear_style_borders(style: Any) -> None:
    """Remove any paragraph border defined on a style (e.g. Title underline)."""
    ppr = style.element.get_or_add_pPr()
    for existing in ppr.findall(qn("w:pBdr")):
        ppr.remove(existing)


def _set_style_bottom_border(style: Any, color: int) -> None:
    """Give a style a thin gray bottom border used as a visual section rule."""
    ppr = style.element.get_or_add_pPr()
    for existing in ppr.findall(qn("w:pBdr")):
        ppr.remove(existing)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")  # eighths of a point -> ~0.5 pt
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), f"{color:06X}")
    pbdr.append(bottom)
    ppr.append(pbdr)


def _apply_document_styles(document: Any) -> None:
    """Define the summary document's typography once, on named styles.

    Sets Normal (body), Title, Heading 1 (section headings), Heading 2 (per-page
    headings), and List Bullet so downstream paragraphs need no per-run
    formatting. Concrete sizes, colors, and spacing live in
    ``autoexcerpter.constants``.
    """
    styles = document.styles

    normal = styles["Normal"]
    _style_font(normal, size_pt=BODY_FONT_PT, bold=False, color=COLOR_BLACK)
    npf = normal.paragraph_format
    npf.space_before = Pt(0)
    npf.space_after = Pt(BODY_SPACE_AFTER_PT)
    npf.line_spacing = 1.0
    npf.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    title = styles["Title"]
    _style_font(title, size_pt=TITLE_FONT_PT, bold=True, color=COLOR_BLACK)
    _clear_style_borders(title)
    tpf = title.paragraph_format
    tpf.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    tpf.space_before = Pt(0)
    tpf.space_after = Pt(TITLE_SPACE_AFTER_PT)

    section_heading = styles["Heading 1"]
    _style_font(
        section_heading,
        size_pt=SECTION_HEADING_FONT_PT,
        bold=True,
        color=COLOR_BLACK,
    )
    _set_style_bottom_border(section_heading, COLOR_SECTION_RULE)
    shp = section_heading.paragraph_format
    shp.space_before = Pt(SECTION_HEADING_SPACE_BEFORE_PT)
    shp.space_after = Pt(SECTION_HEADING_SPACE_AFTER_PT)
    shp.keep_with_next = True

    page_heading = styles["Heading 2"]
    _style_font(
        page_heading,
        size_pt=PAGE_HEADING_FONT_PT,
        bold=True,
        color=COLOR_PAGE_HEADING,
    )
    php = page_heading.paragraph_format
    php.space_before = Pt(PAGE_HEADING_SPACE_BEFORE_PT)
    php.space_after = Pt(PAGE_HEADING_SPACE_AFTER_PT)
    php.keep_with_next = True

    bullet = styles["List Bullet"]
    _style_font(bullet, size_pt=BULLET_FONT_PT, color=COLOR_BLACK)
    bpf = bullet.paragraph_format
    bpf.left_indent = Cm(BULLET_LEFT_INDENT_CM)
    bpf.first_line_indent = -Cm(BULLET_HANGING_INDENT_CM)
    bpf.space_before = Pt(0)
    bpf.space_after = Pt(BULLET_SPACE_AFTER_PT)


def _add_page_number_footer(section: Any) -> None:
    """Add a bottom-right page-number field to a section footer."""
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = paragraph.add_run()
    run.font.name = BODY_FONT_NAME
    run.font.size = Pt(FOOTER_FONT_PT)
    _set_east_asian_font(run._r, BODY_FONT_NAME)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _configure_page(document: Any) -> None:
    """Apply A4 geometry, uniform margins, and a page-number footer."""
    for section in document.sections:
        section.page_width = Cm(PAGE_WIDTH_CM)
        section.page_height = Cm(PAGE_HEIGHT_CM)
        section.top_margin = Cm(PAGE_MARGIN_CM)
        section.bottom_margin = Cm(PAGE_MARGIN_CM)
        section.left_margin = Cm(PAGE_MARGIN_CM)
        section.right_margin = Cm(PAGE_MARGIN_CM)
        _add_page_number_footer(section)


def create_docx_summary(
    summary_results: list[dict[str, Any]],
    output_path: Path,
    document_name: str,
    citation_manager: CitationManager | None = None,
    data: Any = None,
) -> None:
    """Create a DOCX summary document from structured summary results.

    Output order:
    1. Title + Metadata
    2. Document Structure (page type overview)
    3. Content Summaries (in document order)
    4. Consolidated References

    When *citation_manager* and *data* are supplied by the pipeline they are
    already consolidated and OpenAlex-enriched, so this writer renders them as-is
    (one enrichment pass per item, shared with the Markdown writer). When called
    standalone they are built (with default OpenAlex options) and enriched here.
    """
    if citation_manager is None or data is None:
        citation_manager = CitationManager()
        data = prepare_summary_data(summary_results, citation_manager)
        citation_manager.consolidate()
        enrich_if_enabled(citation_manager)

    document = Document()
    _apply_document_styles(document)
    _configure_page(document)
    style_ids = _resolve_style_ids(document)

    _write_title(document, document_name, data, style_ids)
    _write_structure(document, data.page_type_pages, style_ids)
    _write_page_summaries(document, data.page_render_items, style_ids)
    _write_references(document, citation_manager, style_ids)
    _save_atomically(document, output_path)
    logger.info("Summary document saved to %s", output_path)


@dataclass(frozen=True)
class _StyleIds:
    """Resolved ids of the paragraph styles the summary uses."""

    title: str | None
    heading1: str | None
    heading2: str | None
    bullet: str | None


def _resolve_style_ids(document: Any) -> _StyleIds:
    """Resolve the paragraph style ids once per document.

    python-docx's name-based style assignment (add_heading or
    add_paragraph(style=...)) rescans the built-in style catalog on every
    paragraph; setting ``paragraph._p.style = id`` directly avoids that
    O(paragraphs x catalog) cost and produces identical document.xml.
    """
    styles = document.styles
    return _StyleIds(
        title=styles.get_style_id("Title", WD_STYLE_TYPE.PARAGRAPH),
        heading1=styles.get_style_id("Heading 1", WD_STYLE_TYPE.PARAGRAPH),
        heading2=styles.get_style_id("Heading 2", WD_STYLE_TYPE.PARAGRAPH),
        bullet=styles.get_style_id("List Bullet", WD_STYLE_TYPE.PARAGRAPH),
    )


def _write_title(
    document: Any, document_name: str, data: SummaryData, style_ids: _StyleIds
) -> None:
    """Write the title and the centered metadata line."""
    title = document.add_paragraph(sanitize_for_xml(document_name))
    title._p.style = style_ids.title
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    metadata = (
        f"Processed {datetime.now():%Y-%m-%d %H:%M}"
        f" | {data.content_page_count} content pages"
        f" | {data.source_page_count} total pages"
    )
    meta_paragraph = document.add_paragraph()
    meta_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    meta_paragraph.paragraph_format.space_after = Pt(TITLE_SPACE_AFTER_PT)
    meta_run = meta_paragraph.add_run(metadata)
    meta_run.font.name = BODY_FONT_NAME
    meta_run.font.size = Pt(METADATA_FONT_PT)
    meta_run.font.color.rgb = _rgb(COLOR_METADATA_GRAY)


def _write_structure(
    document: Any,
    page_type_pages: dict[str, list[tuple[int, str]]],
    style_ids: _StyleIds,
) -> None:
    """Write the Document Structure section when any page type has pages."""
    if not any(pages for pages in page_type_pages.values()):
        return
    structure_heading = document.add_paragraph("Document Structure")
    structure_heading._p.style = style_ids.heading1

    for pt in STRUCTURE_PAGE_TYPE_ORDER:
        pages = page_type_pages.get(pt, [])
        if pages:
            label = PAGE_TYPE_LABELS.get(pt, pt.replace("_", " ").title())
            page_range = format_structure_page_range(pages)
            struct_para = document.add_paragraph()
            label_run = struct_para.add_run(f"{label}: ")
            label_run.bold = True
            struct_para.add_run(page_range)


def _write_page_summaries(
    document: Any, page_render_items: list[PageRenderData], style_ids: _StyleIds
) -> None:
    """Write one heading and its bullet points per summarized page."""
    for page_item in page_render_items:
        # The heading is the one string reaching add_paragraph unsanitized: a
        # lenient endpoint can return a page_number carrying an XML-illegal
        # control character, which would fail the entire save.
        page_heading = document.add_paragraph(sanitize_for_xml(page_item.heading_text))
        page_heading._p.style = style_ids.heading2

        for point in page_item.bullet_points:
            paragraph = document.add_paragraph()
            paragraph._p.style = style_ids.bullet
            add_formatted_text_to_paragraph(
                paragraph, collapse_internal_newlines(point)
            )


def _write_references(
    document: Any, citation_manager: CitationManager, style_ids: _StyleIds
) -> None:
    """Write the Consolidated References section on a new page."""
    if not citation_manager.citations:
        return
    logger.info(
        "Processing %d unique citations for consolidated references section",
        len(citation_manager.citations),
    )

    document.add_page_break()
    references_heading = document.add_paragraph("Consolidated References")
    references_heading._p.style = style_ids.heading1

    note_text = (
        "The following references were extracted from the document and "
        "consolidated. Duplicate citations have been merged, showing all "
        "pages where each citation appears. Where available, hyperlinks "
        "provide access to extended metadata via OpenAlex."
    )
    note_paragraph = document.add_paragraph()
    note_run = note_paragraph.add_run(note_text)
    note_run.italic = True
    note_run.font.size = Pt(METADATA_FONT_PT)
    note_run.font.color.rgb = _rgb(COLOR_METADATA_GRAY)

    citations_with_pages = citation_manager.get_citations_with_pages()
    for idx, (citation, page_range_str) in enumerate(citations_with_pages, start=1):
        _write_reference(document, idx, citation, page_range_str)


def _write_reference(
    document: Any, idx: int, citation: Citation, page_range_str: str
) -> None:
    """Write one numbered reference with its pages and OpenAlex metadata."""
    ref_paragraph = document.add_paragraph()
    rpf = ref_paragraph.paragraph_format
    rpf.left_indent = Cm(REF_HANGING_INDENT_CM)
    rpf.first_line_indent = -Cm(REF_HANGING_INDENT_CM)
    rpf.space_after = Pt(REF_SPACE_AFTER_PT)

    num_run = ref_paragraph.add_run(f"[{idx}] ")
    num_run.bold = True
    num_run.font.size = Pt(REF_FONT_PT)

    citation_text = sanitize_for_xml(collapse_internal_newlines(citation.raw_text))

    if citation.url:
        # A hyperlink is a single styled run; drop emphasis markers.
        add_hyperlink(
            ref_paragraph,
            citation.url,
            strip_markdown_emphasis(citation_text),
            size_pt=REF_FONT_PT,
        )
    else:
        for kind, content in parse_markdown_emphasis(citation_text):
            text_run = ref_paragraph.add_run(content)
            text_run.font.size = Pt(REF_FONT_PT)
            if kind in ("italic", "bold_italic"):
                text_run.italic = True
            if kind in ("bold", "bold_italic"):
                text_run.bold = True

    if page_range_str:
        page_run = ref_paragraph.add_run(f" ({page_range_str})")
        page_run.italic = True
        page_run.font.size = Pt(REF_FONT_PT)

    if citation.metadata:
        meta_info_parts = []
        if citation.doi:
            meta_info_parts.append(f"DOI: {citation.doi}")
        if citation.metadata.get("publication_year"):
            meta_info_parts.append(f"Year: {citation.metadata['publication_year']}")

        if meta_info_parts:
            meta_run = ref_paragraph.add_run(f" [{', '.join(meta_info_parts)}]")
            meta_run.font.size = Pt(REF_META_FONT_PT)
            meta_run.italic = True
            meta_run.font.color.rgb = _rgb(COLOR_REF_META_GRAY)


def _save_atomically(document: Any, output_path: Path) -> None:
    """Save to a sibling ``.tmp`` file and replace the target atomically.

    A crash mid-save never leaves a truncated .docx that resume would trust as
    complete (resume classifies on ``exists()`` and ``st_size > 0``); the
    temp file is removed on any failure.
    """
    tmp_path = output_path.with_name(output_path.name + ".tmp")
    try:
        document.save(str(tmp_path))
        os.replace(tmp_path, output_path)
    except Exception:
        with contextlib.suppress(OSError):
            tmp_path.unlink()
        raise
