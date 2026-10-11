"""Import the outputs of version 1.0 runs into the store (``migrate``).

:func:`migrate` finds, at or below its input path:

- the ``autoexcerpter.sqlite`` of every output folder;
- the working logs ``transcription.jsonl`` and ``summary.jsonl`` that
  version 1.0 kept per item in ``<stem>_autoexcerpter/``;
- the rendered files of each item (``<stem>_transcription.md`` or ``.txt``,
  ``<stem>_summary.md`` and ``<stem>_summary.docx``).

Every database is checked before anything is written: a newer schema, an
unreadable file and a versionless layout other than the tables of version
1.0 (``documents``, ``pages``, ``summaries`` and ``citations``) raise the
store's :class:`~autoexcerpter.common.store.StoreError` (exit code 2). A
database is snapshotted through the SQLite backup API before its first
write, as ``autoexcerpter.pre-migrate-<time>.sqlite`` beside it. A version
1.0 database is rebuilt into the store schema in one transaction
(:func:`convert_versionless_ae`), its row counts verified against the
snapshot; the pages of a document's ``pages_total`` that its rows lack are
planned ``pending``. Naive version 1.0 timestamps are read as this
machine's local time and stored in UTC.

Per item the import writes, under the document's lease and in one
transaction that first checks that no run changed the planned units since
the plan:

- the last record of every page in the transcription log as its
  transcription unit and the last record of the summary log as its summary
  unit and ``summaries`` row, with the conventions of
  :mod:`autoexcerpter.output.store`; the pages the logs lack are planned
  ``pending``, so a partial item resumes where version 1.0 would have;
- the image-settings record of the log header, the input path and hash;
- the citations, consolidated offline from the summaries' references with
  metadata from the OpenAlex cache only, when the store holds none;
- the existing rendered files: one that holds its rendering from the store
  with its hash, so that ``render`` replaces it in place; one that differs
  or cannot be read without a hash, so that ``render`` keeps it unless
  forced.

A unit that a run of the store wrote is kept as it is. A rerun imports
nothing and changes no row; old files are only read. Each document is then
verified from the store: unit counts and failures, and the re-rendered
bodies, the text formats compared without the header, the
``<page_break .../>`` markers and one missing final newline, the Word
summary by its paragraph, table and equation text without the metadata
line (:func:`docx_text`). A difference or an
unreadable file is a mismatch (exit code 1). The report also records each
document's resume state before and after the import.

Version 1.0 working logs carry a header with ``_format_version`` 2; every
following line is one page result keyed by ``original_input_order_index``.
:func:`read_working_log` reads one such log tolerantly: a truncated last
line (a crash mid-append) is dropped, unparsable interior lines and records
without a page index are skipped, and a log without the version 2 header is
refused.
"""

from __future__ import annotations

import io
import json
import logging
import re
import sqlite3
import uuid
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Final

from docx import Document
from docx.oxml.ns import qn

from autoexcerpter.common.sqlite import read, snapshot
from autoexcerpter.common.store import (
    ConversionResult,
    Converter,
    PlannedUnit,
    SchemaState,
    Store,
    StoreError,
    UnrecognizedSchemaError,
    check_schema,
    convert_versionless,
    file_sha256,
    inspect_schema,
    refresh_documents,
)
from autoexcerpter.constants import is_blank_transcription
from autoexcerpter.llm.capabilities import detect_provider
from autoexcerpter.output.paths import (
    LEGACY_LOGS,
    LEGACY_SUFFIX,
    ItemPaths,
    item_base,
)
from autoexcerpter.output.render import (
    RENDERERS,
    SUMMARY_DOCX,
    SUMMARY_MD,
    TRANSCRIPTION_MD,
    TRANSCRIPTION_TXT,
)
from autoexcerpter.output.store import (
    COMPONENT,
    DATABASE_NAME,
    KIND_SUMMARY,
    KIND_TRANSCRIPTION,
    TOOL,
    CitationRow,
    InputMismatchError,
    ResumeRequest,
    SummaryRow,
    classify,
    database_path,
    document_kind,
    image_settings_record,
    open_store,
    transcription_outcome,
)
from autoexcerpter.pipeline.pages import summary_row
from autoexcerpter.rendering.citations import OpenAlexCache, OpenAlexClient
from autoexcerpter.rendering.page_numbering import PageNumberProcessor
from autoexcerpter.rendering.summary import build_render_context

__all__ = [
    "CONVERTED_COMMAND",
    "FORMAT_VERSION",
    "MARKER",
    "MIGRATE_COMMAND",
    "PAGE_KEY",
    "VERSION_KEY",
    "DatabaseReport",
    "DocumentReport",
    "LegacyItem",
    "MigrationError",
    "MigrationReport",
    "WorkingLog",
    "WorkingLogError",
    "body_matches",
    "body_problem",
    "convert_versionless_ae",
    "docx_text",
    "find_items",
    "migrate",
    "offline_citations",
    "read_working_log",
    "snapshot_path",
]

logger = logging.getLogger(__name__)

VERSION_KEY: Final = "_format_version"
FORMAT_VERSION: Final = 2
PAGE_KEY: Final = "original_input_order_index"

MIGRATE_COMMAND: Final = "migrate"
CONVERTED_COMMAND: Final = "migrate (converted)"
"""The command of the runs rows a versionless conversion creates."""

SNAPSHOT_PREFIX: Final = "autoexcerpter.pre-migrate-"

MARKER: Final = re.compile(r"^<page_break (?:pdf|image)=\"\d+\"/>\n", re.MULTILINE)
"""A page-break marker line, as version 1.0 and ``add_page_breaks.py`` write it."""

_SUMMARY_LINE: Final = re.compile(r"^\*Summary generated .*\*\n", re.MULTILINE)
_DOCX_METADATA: Final = re.compile(r"^Processed .* content pages \| \d+ total pages$")
_W_P: Final = qn("w:p")
_W_T: Final = qn("w:t")
_M_MATH: Final = frozenset({qn("m:oMath"), qn("m:oMathPara")})

_V1_COLUMNS: Final[Mapping[str, frozenset[str]]] = {
    "documents": frozenset(
        {"name", "path", "kind", "hash", "transcription_model", "summary_model"}
        | {"spec", "outputs", "updated_at"}
    ),
    "pages": frozenset(
        {"document", "page_index", "image", "pdf_page_index", "status", "error"}
        | {"model", "usage", "text"}
    ),
    "summaries": frozenset(
        {"document", "page_index", "status", "error", "model", "usage"}
        | {"page_information", "bullet_points", "page_references"}
    ),
    "citations": frozenset(
        {"document", "position", "text", "pages", "partial", "doi", "url", "metadata"}
    ),
}

_UPSERT_UNIT: Final = """
INSERT INTO units (
    doc_id, kind, idx, status, outcome, source_index, label, content, tool,
    model, provider, run_id, attempts, usage, error, error_type, provenance,
    stale, updated_at
) VALUES (
    :doc_id, :kind, :idx, :status, :outcome, :source_index, :label, :content,
    :tool, :model, :provider, :run_id, :attempts, :usage, :error, :error_type,
    :provenance, 0, :now
)
ON CONFLICT (doc_id, kind, idx) DO UPDATE SET
    status = excluded.status, outcome = excluded.outcome,
    source_index = excluded.source_index, label = excluded.label,
    content = excluded.content, tool = excluded.tool, model = excluded.model,
    provider = excluded.provider, run_id = excluded.run_id,
    attempts = excluded.attempts, usage = excluded.usage,
    error = excluded.error, error_type = excluded.error_type,
    provenance = excluded.provenance, response_id = NULL, stale = 0,
    updated_at = excluded.updated_at
"""
_UPSERT_SUMMARY: Final = """
INSERT INTO summaries (doc_id, idx, page, page_information, bullet_points,
    page_references) VALUES (?, ?, ?, ?, ?, ?)
ON CONFLICT (doc_id, idx) DO UPDATE SET page = excluded.page,
    page_information = excluded.page_information,
    bullet_points = excluded.bullet_points,
    page_references = excluded.page_references
"""
_PLAN_UNIT: Final = """
INSERT INTO units (doc_id, kind, idx, status, source_index, label, tool, model,
    provider, run_id, attempts, stale, updated_at)
VALUES (?, ?, ?, 'pending', ?, ?, ?, ?, ?, ?, 0, 0, ?)
ON CONFLICT (doc_id, kind, idx) DO NOTHING
"""
_REGISTER_OUTPUT: Final = """
INSERT INTO outputs (doc_id, path, format, sha256, rendered_at, run_id, stale)
VALUES (?, ?, ?, ?, ?, ?, 0)
ON CONFLICT (doc_id, path) DO NOTHING
"""


class MigrationError(Exception):
    """A conversion whose verification failed; exit code 1."""

    exit_code: int = 1


# =============================================================================
# Working logs
# =============================================================================


class WorkingLogError(ValueError):
    """A working log lacks the version 2 header."""


@dataclass(frozen=True)
class WorkingLog:
    """A version 1.0 working log: its header and its page results in file order.

    *dropped_tail* says whether a truncated last line was dropped;
    *skipped_lines* holds the 1-based numbers of unparsable interior lines.
    """

    header: dict[str, Any]
    results: list[dict[str, Any]]
    dropped_tail: bool = False
    skipped_lines: tuple[int, ...] = field(default=())


def read_working_log(path: Path) -> WorkingLog | None:
    """Read the working log at *path*; None when it is missing or blank.

    Raises:
        WorkingLogError: The first line is not a version 2 header.
        OSError: The file cannot be read.
        UnicodeDecodeError: The file is not UTF-8.
    """
    if not path.exists():
        return None
    text = path.read_text(encoding="utf-8").removeprefix("\N{BYTE ORDER MARK}")
    lines = [
        (number, line.strip())
        for number, line in enumerate(text.split("\n"), start=1)
        if line.strip()
    ]
    if not lines:
        return None
    try:
        header = json.loads(lines[0][1])
    except json.JSONDecodeError as exc:
        raise WorkingLogError(f"{path}: the first line is not a header: {exc}") from exc
    if not isinstance(header, dict) or header.get(VERSION_KEY) != FORMAT_VERSION:
        found = header.get(VERSION_KEY) if isinstance(header, dict) else None
        raise WorkingLogError(
            f"{path}: expected {VERSION_KEY}={FORMAT_VERSION} in the header, "
            f"found {found!r}."
        )
    results: list[dict[str, Any]] = []
    skipped: list[int] = []
    dropped_tail = False
    last = len(lines) - 1
    for position, (number, line) in enumerate(lines[1:], start=1):
        try:
            record = json.loads(line)
        except json.JSONDecodeError:
            if position == last:
                dropped_tail = True
            else:
                skipped.append(number)
            continue
        if isinstance(record, dict) and isinstance(record.get(PAGE_KEY), int):
            results.append(record)
    if skipped:
        logger.warning("%s: skipped unparsable line(s) %s", path, skipped)
    return WorkingLog(header, results, dropped_tail, tuple(skipped))


# =============================================================================
# Values
# =============================================================================


def _now() -> str:
    return datetime.now(UTC).isoformat(timespec="milliseconds")


def _utc(stamp: Any) -> str:
    """Return a version 1.0 timestamp in UTC; now when it is missing or invalid.

    Version 1.0 wrote naive local times, which are read as this machine's
    local time.
    """
    try:
        moment = datetime.fromisoformat(str(stamp)) if stamp else None
    except ValueError:
        moment = None
    if moment is None:
        return _now()
    if moment.tzinfo is None:
        moment = moment.astimezone()
    return moment.astimezone(UTC).isoformat(timespec="milliseconds")


def _plain(value: Any) -> Any:
    """Return *value* as stored JSON reads back (tuples become lists)."""
    if value is None:
        return None
    return json.loads(json.dumps(value, ensure_ascii=False, default=str))


def _mapping(value: Any) -> dict[str, Any]:
    return dict(value) if isinstance(value, Mapping) else {}


def _int(value: Any) -> int | None:
    return value if isinstance(value, int) and not isinstance(value, bool) else None


def _text(value: Any) -> str | None:
    return value if isinstance(value, str) and value else None


def _json(value: Any) -> str | None:
    if value is None:
        return None
    return json.dumps(value, ensure_ascii=False, default=str)


def _json_value(value: Any) -> Any:
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value
    return value


def _ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _rows(
    connection: sqlite3.Connection, sql: str, params: Sequence[Any] = ()
) -> list[dict[str, Any]]:
    cursor = connection.execute(sql, tuple(params))
    names = [column[0] for column in cursor.description]
    return [dict(zip(names, row, strict=True)) for row in cursor.fetchall()]


def _provider(model: str, given: Any = None) -> str:
    """Return the recorded provider, else the one the model id names."""
    return _text(given) or (detect_provider(model) if model else "unknown")


def summary_outcome(
    text: str | None, error: str | None, transcription_failed: bool
) -> str:
    """Return the outcome of a page's summary unit (see the store module)."""
    if transcription_failed:
        return "placeholder"
    if error is not None:
        return "failed"
    return "blank" if is_blank_transcription(text) else "ok"


def snapshot_path(database: Path) -> Path:
    """Return a free snapshot path beside *database* for this moment."""
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%S")
    candidate = database.with_name(f"{SNAPSHOT_PREFIX}{stamp}.sqlite")
    number = 1
    while candidate.exists():
        number += 1
        candidate = database.with_name(f"{SNAPSHOT_PREFIX}{stamp}-{number}.sqlite")
    return candidate


def _output_key(path: Path, root: Path) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(root.resolve()).as_posix()
    except ValueError:
        return str(resolved)


def item_files(name: str, root: Path) -> list[tuple[Path, str]]:
    """Return the rendered files of item *name* in *root* that exist."""
    md = ItemPaths.for_item(name, root, "md")
    txt = ItemPaths.for_item(name, root, "txt")
    files = (
        (md.transcription, TRANSCRIPTION_MD),
        (txt.transcription, TRANSCRIPTION_TXT),
        (md.summary_md, SUMMARY_MD),
        (md.summary_docx, SUMMARY_DOCX),
    )
    return [(path, fmt) for path, fmt in files if path.is_file()]


# =============================================================================
# Bodies
# =============================================================================


def _body(fmt: str, text: str) -> str:
    if fmt == SUMMARY_MD:
        text = _SUMMARY_LINE.sub("", text, count=1)
    else:
        text = text.partition("\n\n")[2]
    return MARKER.sub("", text)


def body_matches(fmt: str, rendered: str, existing: str) -> bool:
    """Whether a rendering has the body of an existing file of format *fmt*.

    The transcription header and the summary's metadata line are left out,
    the page-break markers are stripped from both, and a missing final
    newline of the existing file is tolerated.
    """
    new, old = _body(fmt, rendered), _body(fmt, existing)
    return new in (old, old + "\n")


def _math(element: Any) -> str:
    """Return an equation as its element names and texts, properties left out."""
    name = element.tag.rpartition("}")[2]
    if name.endswith("Pr"):
        return ""
    text = (element.text or "") if name == "t" else ""
    inner = "".join(_math(child) for child in element)
    return f"<{name}>{text}{inner}</{name}>"


def _paragraph_text(element: Any) -> str:
    """Return a paragraph's run texts and equations in document order."""
    parts = []
    for child in element:
        if child.tag == _W_T:
            parts.append(child.text or "")
        elif child.tag in _M_MATH:
            parts.append(_math(child))
        else:
            parts.append(_paragraph_text(child))
    return "".join(parts)


def docx_text(data: bytes) -> list[str]:
    """Return the content of a Word summary, normalized for comparison.

    One entry per non-empty paragraph of the body in document order, table
    cells included, stripped: its texts and its equations (as their OMML
    element names and texts, properties left out). The metadata line
    (``Processed ... | N content pages | M total pages``) is left out.

    Raises:
        Exception: *data* is not a Word document python-docx can read.
    """
    document = Document(io.BytesIO(data))
    lines = []
    for paragraph in document.element.body.iter(_W_P):
        text = _paragraph_text(paragraph).strip()
        if text and not _DOCX_METADATA.match(text):
            lines.append(text)
    return lines


def body_problem(
    connection: sqlite3.Connection, doc_id: int, path: Path, fmt: str
) -> str | None:
    """Return why the file at *path* does not hold its rendering, or None.

    The rendering comes from *connection*; the text formats are compared by
    :func:`body_matches`, the Word summary by :func:`docx_text`. A file that
    cannot be read, or a Word summary that cannot be compared, is a problem.
    """
    try:
        if fmt == SUMMARY_DOCX:
            existing: str | list[str] = docx_text(path.read_bytes())
        else:
            existing = path.read_text(encoding="utf-8")
    except Exception as exc:  # noqa: BLE001 - any unreadable file is reported
        return f"{path.name} could not be read: {exc}"
    rendered = RENDERERS[fmt](connection, doc_id)
    if isinstance(existing, list):
        data = rendered if isinstance(rendered, bytes) else rendered.encode("utf-8")
        try:
            same = docx_text(data) == existing
        except Exception as exc:  # noqa: BLE001 - an unverified file stays unhashed
            return f"{path.name} could not be compared: {exc}"
    else:
        same = isinstance(rendered, str) and body_matches(fmt, rendered, existing)
    if same:
        return None
    return (
        f"{path.name} differs from its rendering from the store; it is "
        "registered without a hash, so render replaces it only with --force."
    )


# =============================================================================
# Offline citations
# =============================================================================


def offline_citations(
    summaries: list[dict[str, Any]], position_label: str, cache: OpenAlexCache
) -> list[CitationRow]:
    """Return the consolidated citations of *summaries* without a lookup.

    The citations are consolidated as an item run does; their metadata comes
    from the OpenAlex *cache* alone, and no request is made.
    """
    adjusted = PageNumberProcessor().adjust_and_sort_page_numbers(summaries)
    manager, _ = build_render_context(
        adjusted,
        position_label=position_label,
        openalex_enabled=True,
        max_requests=0,
        openalex=OpenAlexClient(cache=cache),
    )
    if manager.citations:
        manager.enrich_with_metadata(max_requests=0)
    return [
        CitationRow(
            text=citation.raw_text,
            pages=pages,
            partial=citation.partial,
            doi=citation.doi,
            url=citation.url,
            metadata=citation.metadata,
        )
        for citation, pages in manager.get_citations_with_pages()
    ]


# =============================================================================
# Conversion of version 1.0 databases
# =============================================================================


def _versionless_matches(tables: Mapping[str, tuple[str, ...]]) -> bool:
    if set(tables) != set(_V1_COLUMNS):
        return False
    return all(set(tables[name]) >= columns for name, columns in _V1_COLUMNS.items())


@dataclass(frozen=True)
class _Page:
    failed: bool
    text: str | None
    label: str | None
    source_index: int | str | None


def _convert_pages(
    connection: sqlite3.Connection,
    table: str,
    doc: Mapping[str, Any],
    context: Mapping[str, Any],
) -> dict[int, _Page]:
    pages: dict[int, _Page] = {}
    fallback = _text(doc.get("transcription_model")) or ""
    for page in _rows(
        connection,
        f"SELECT * FROM {_ident(table)} WHERE document = ? ORDER BY page_index",
        (doc["name"],),
    ):
        idx = int(page["page_index"])
        failed = page.get("status") == "failed"
        text = page.get("text")
        label = _text(page.get("image"))
        pdf_index = _int(page.get("pdf_page_index"))
        source: int | str | None = pdf_index
        if source is None:
            source = idx if context["kind"] == "pdf" else label
        error = (_text(page.get("error")) or "failed") if failed else None
        model = _text(page.get("model")) or fallback
        pages[idx] = _Page(failed, text, label, source)
        connection.execute(
            _UPSERT_UNIT,
            {
                **context["unit"],
                "kind": KIND_TRANSCRIPTION,
                "idx": idx,
                "status": "failed" if failed else "done",
                "outcome": transcription_outcome(text, error),
                "source_index": source,
                "label": label,
                "content": text,
                "model": model,
                "provider": _provider(model),
                "attempts": 1,
                "usage": page.get("usage"),
                "error": error,
                "error_type": page.get("error_type") if failed else None,
                "provenance": None,
            },
        )
    return pages


def _convert_summaries(
    connection: sqlite3.Connection,
    table: str,
    doc: Mapping[str, Any],
    context: Mapping[str, Any],
    pages: Mapping[int, _Page],
) -> None:
    fallback = _text(doc.get("summary_model")) or ""
    for summary in _rows(
        connection,
        f"SELECT * FROM {_ident(table)} WHERE document = ? ORDER BY page_index",
        (doc["name"],),
    ):
        idx = int(summary["page_index"])
        page = pages.get(idx, _Page(False, None, None, None))
        failed = summary.get("status") == "failed" and not page.failed
        error = (_text(summary.get("error")) or "failed") if failed else None
        model = _text(summary.get("model")) or fallback
        connection.execute(
            _UPSERT_UNIT,
            {
                **context["unit"],
                "kind": KIND_SUMMARY,
                "idx": idx,
                "status": "failed" if failed else "done",
                "outcome": summary_outcome(page.text, error, page.failed),
                "source_index": page.source_index,
                "label": page.label,
                "content": None,
                "model": model,
                "provider": _provider(model),
                "attempts": 1,
                "usage": summary.get("usage"),
                "error": error,
                "error_type": None,
                "provenance": None,
            },
        )
        connection.execute(
            _UPSERT_SUMMARY,
            (
                context["unit"]["doc_id"],
                idx,
                idx + 1,
                summary.get("page_information"),
                summary.get("bullet_points"),
                summary.get("page_references"),
            ),
        )


def _convert_document(
    connection: sqlite3.Connection,
    legacy: Mapping[str, str],
    doc: Mapping[str, Any],
    run_id: str,
) -> None:
    name = str(doc["name"])
    stamp = _utc(doc.get("updated_at"))
    kind = document_kind(str(doc.get("kind") or ""))
    (doc_id,) = connection.execute(
        "INSERT INTO documents (name, tool, input_path, input_sha256, kind, "
        "first_run_id, last_run_id, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?) "
        "RETURNING doc_id",
        (name, TOOL, doc.get("path"), doc.get("hash"), kind, run_id, run_id, stamp),
    ).fetchone()
    context = {
        "kind": kind,
        "unit": {"doc_id": doc_id, "tool": TOOL, "run_id": run_id, "now": stamp},
    }
    pages = _convert_pages(connection, legacy["pages"], doc, context)
    _convert_summaries(connection, legacy["summaries"], doc, context, pages)
    _plan_missing_pages(connection, doc, context)
    connection.execute(
        "INSERT INTO citations (doc_id, position, text, pages, partial, doi, url, "
        f"metadata) SELECT ?, position, text, pages, partial, doi, url, metadata "
        f"FROM {_ident(legacy['citations'])} WHERE document = ?",
        (doc_id, name),
    )


def _plan_missing_pages(
    connection: sqlite3.Connection, doc: Mapping[str, Any], context: Mapping[str, Any]
) -> None:
    """Plan the pages of ``pages_total`` the version 1.0 rows lack as pending.

    Summary units are planned when the document recorded a summary model.
    """
    total = _int(doc.get("pages_total"))
    if not total:
        return
    unit = context["unit"]
    models = {KIND_TRANSCRIPTION: _text(doc.get("transcription_model"))}
    if _text(doc.get("summary_model")):
        models[KIND_SUMMARY] = _text(doc.get("summary_model"))
    for kind, model in models.items():
        name = model or ""
        connection.executemany(
            _PLAN_UNIT,
            [
                (unit["doc_id"], kind, idx, idx if context["kind"] == "pdf" else None)
                + (None, TOOL, name, _provider(name), unit["run_id"], unit["now"])
                for idx in range(total)
            ],
        )


def _register_output(
    connection: sqlite3.Connection,
    doc_id: int,
    path: Path,
    fmt: str,
    root: Path,
    run_id: str,
) -> None:
    """Register an existing file; with its hash only when it holds its rendering.

    A file that differs from its rendering, or cannot be read, is registered
    without a hash, so that ``render`` keeps it unless forced.
    """
    digest = None
    if body_problem(connection, doc_id, path, fmt) is None:
        digest = file_sha256(path)
    connection.execute(
        _REGISTER_OUTPUT,
        (doc_id, _output_key(path, root), fmt, digest, _now(), run_id),
    )


def _rebuild_versionless_ae(
    connection: sqlite3.Connection, legacy: Mapping[str, str]
) -> None:
    """Copy the version 1.0 rows into the store tables.

    One ``runs`` row stands for each distinct effective spec and tool
    version the documents recorded; its command is :data:`CONVERTED_COMMAND`.
    The existing files are registered by the import that follows, once the
    document holds everything it renders from.
    """
    documents = _rows(
        connection, f"SELECT * FROM {_ident(legacy['documents'])} ORDER BY name"
    )
    runs: dict[tuple[str, str], str] = {}
    for doc in documents:
        spec = _text(doc.get("spec"))
        version = _text(doc.get("tool_version")) or "unknown"
        run_id = runs.get((spec or "", version))
        if run_id is None:
            run_id = str(uuid.uuid4())
            runs[(spec or "", version)] = run_id
            stamp = _utc(doc.get("updated_at"))
            connection.execute(
                "INSERT INTO runs (run_id, tool, tool_version, command, started_at, "
                "ended_at, spec, settings_fingerprint) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    run_id,
                    TOOL,
                    version,
                    CONVERTED_COMMAND,
                    stamp,
                    stamp,
                    spec,
                    doc.get("settings_fingerprint"),
                ),
            )
        _convert_document(connection, legacy, doc, run_id)
    refresh_documents(connection)


def _counts(connection: sqlite3.Connection, store: bool) -> dict[str, Any]:
    if store:
        sql = (
            "SELECT d.name AS document, u.kind AS kind, count(*) AS total, "
            "sum(u.status = 'failed') AS failed FROM units AS u JOIN documents AS d "
            "ON d.doc_id = u.doc_id WHERE u.status != 'pending' GROUP BY d.name, u.kind"
        )
        citations = (
            "SELECT d.name, count(*) FROM citations AS c JOIN documents AS d "
            "ON d.doc_id = c.doc_id GROUP BY d.name"
        )
    else:
        sql = (
            "SELECT document, 'transcription' AS kind, count(*) AS total, "
            "sum(status = 'failed') AS failed FROM pages GROUP BY document "
            "UNION ALL SELECT s.document, 'summary', count(*), "
            "sum(s.status = 'failed' AND coalesce(p.status, '') != 'failed') "
            "FROM summaries AS s LEFT JOIN pages AS p ON p.document = s.document "
            "AND p.page_index = s.page_index GROUP BY s.document"
        )
        citations = "SELECT document, count(*) FROM citations GROUP BY document"
    return {
        "documents": connection.execute("SELECT count(*) FROM documents").fetchone()[0],
        "units": {
            (row["document"], row["kind"]): (row["total"], row["failed"])
            for row in _rows(connection, sql)
        },
        "citations": dict(connection.execute(citations).fetchall()),
    }


def _expected_coverage(connection: sqlite3.Connection) -> dict[tuple[str, str], int]:
    """Return the units each version 1.0 document needs, per kind.

    A document needs its stored pages and the first ``pages_total`` pages as
    transcription units, and as summary units its stored summaries and,
    when it recorded a summary model, the first ``pages_total`` pages.
    """
    stored: dict[tuple[str, str], set[int]] = {}
    for table, kind in (("pages", KIND_TRANSCRIPTION), ("summaries", KIND_SUMMARY)):
        for document, idx in connection.execute(
            f"SELECT document, page_index FROM {table}"
        ):
            stored.setdefault((str(document), kind), set()).add(int(idx))
    coverage: dict[tuple[str, str], int] = {}
    for doc in _rows(connection, "SELECT * FROM documents"):
        name, total = str(doc["name"]), _int(doc.get("pages_total")) or 0
        summarized = bool(_text(doc.get("summary_model")))
        for kind, needed in (
            (KIND_TRANSCRIPTION, total),
            (KIND_SUMMARY, total if summarized else 0),
        ):
            count = len(stored.get((name, kind), set()) | set(range(needed)))
            if count:
                coverage[(name, kind)] = count
    return coverage


def _coverage(connection: sqlite3.Connection) -> dict[tuple[str, str], int]:
    return {
        (str(name), str(kind)): int(count)
        for name, kind, count in connection.execute(
            "SELECT d.name, u.kind, count(*) FROM units AS u JOIN documents AS d "
            "ON d.doc_id = u.doc_id GROUP BY d.name, u.kind"
        )
    }


def _verify_versionless_ae(connection: sqlite3.Connection, saved: Path) -> None:
    """Compare the converted rows with the snapshot of the old tables.

    The done and failed units, their failures and the citations must equal
    the old rows; with the pending units, every document must hold the
    units its ``pages_total`` needs.
    """

    def before(snapshot_connection: sqlite3.Connection) -> dict[str, Any]:
        counts = _counts(snapshot_connection, False)
        counts["coverage"] = _expected_coverage(snapshot_connection)
        return counts

    old = read(saved, before)
    new = _counts(connection, True)
    new["coverage"] = _coverage(connection)
    if old != new:
        raise MigrationError(
            f"The converted database does not hold the rows of {saved.name}: "
            f"{old['documents']} documents before, {new['documents']} after; "
            f"units per document {old['units']} before, {new['units']} after; "
            f"citations {old['citations']} before, {new['citations']} after; "
            f"expected units {old['coverage']}, stored {new['coverage']}."
        )


def convert_versionless_ae(database: Path, saved: Path) -> ConversionResult:
    """Convert a version 1.0 AutoExcerpter database into the store schema.

    The database is snapshotted to *saved* first; the conversion runs in one
    transaction and is verified against the snapshot. A store database is
    left alone (status ``current``).

    Raises:
        UnrecognizedSchemaError: Another versionless layout.
        MigrationError: The converted row counts differ; nothing changed.
    """
    converter = Converter(
        "autoexcerpter-1.0",
        _versionless_matches,
        _rebuild_versionless_ae,
        _verify_versionless_ae,
    )
    return convert_versionless(
        database, [converter], snapshot_path=saved, components=(COMPONENT,)
    )


def _check_database(path: Path) -> SchemaState:
    """Return the schema state of *path*; raise the store's error to refuse it."""
    state = inspect_schema(path)
    if state.versionless:
        if not _versionless_matches(state.tables):
            raise UnrecognizedSchemaError(
                f"{path} holds tables {sorted(state.tables)} that migrate does not "
                "recognize; it is left as it is."
            )
        return state
    return check_schema(path, (COMPONENT,))


# =============================================================================
# Legacy items
# =============================================================================


@dataclass(frozen=True)
class LegacyItem:
    """The working folder of one version 1.0 item and the folder it renders in."""

    folder: Path

    @property
    def root(self) -> Path:
        """Return the output folder of the item."""
        return self.folder.parent

    @property
    def base(self) -> str:
        """Return the file-name base of the item's outputs."""
        return self.folder.name.removesuffix(LEGACY_SUFFIX)

    @property
    def logs(self) -> tuple[Path, Path]:
        """Return the transcription and the summary log."""
        transcription, summary = LEGACY_LOGS
        return self.folder / transcription, self.folder / summary


def _is_working_folder(folder: Path) -> bool:
    return (
        folder.is_dir()
        and folder.name.endswith(LEGACY_SUFFIX)
        and folder.name != LEGACY_SUFFIX
        and any((folder / log).is_file() for log in LEGACY_LOGS)
    )


def find_items(path: Path) -> list[LegacyItem]:
    """Return the version 1.0 working folders at or below *path*.

    A log file selects its folder; a database selects nothing.
    """
    path = Path(path).resolve()
    if path.is_file():
        folder = path.parent
        if path.name in LEGACY_LOGS and _is_working_folder(folder):
            return [LegacyItem(folder)]
        return []
    folders = [path] if _is_working_folder(path) else []
    folders += sorted(
        found for found in path.rglob(f"*{LEGACY_SUFFIX}") if _is_working_folder(found)
    )
    return [LegacyItem(folder) for folder in folders]


@dataclass
class _Unit:
    """A unit as the import writes it."""

    kind: str
    idx: int
    status: str
    model: str
    provider: str
    content: str | None = None
    outcome: str | None = None
    label: str | None = None
    source_index: int | str | None = None
    usage: Any = None
    error: str | None = None
    error_type: str | None = None
    provenance: Any = None
    attempts: int = 1

    COMPARED = (
        "status",
        "outcome",
        "content",
        "label",
        "source_index",
        "model",
        "provider",
        "usage",
        "error",
        "error_type",
        "provenance",
        "attempts",
    )

    def differs(self, row: Mapping[str, Any]) -> bool:
        """Whether the stored *row* holds other values than this unit."""
        for name in self.COMPARED:
            stored = row.get(name)
            if name in ("usage", "provenance"):
                stored = _json_value(stored)
            if getattr(self, name) != stored:
                return True
        return False

    def params(self, doc_id: int, run_id: str, now: str) -> dict[str, Any]:
        """Return the parameters of the unit's upsert."""
        return {
            "doc_id": doc_id,
            "kind": self.kind,
            "idx": self.idx,
            "status": self.status,
            "outcome": self.outcome,
            "source_index": self.source_index,
            "label": self.label,
            "content": self.content,
            "tool": TOOL,
            "model": self.model,
            "provider": self.provider,
            "run_id": run_id,
            "attempts": self.attempts,
            "usage": _json(self.usage),
            "error": self.error,
            "error_type": self.error_type,
            "provenance": _json(self.provenance),
            "now": now,
        }


@dataclass
class LegacyDocument:
    """What the store should hold for one version 1.0 item after the import.

    *pending* holds the pages the logs lack, per unit kind, with the model
    and provider of that kind in *venues*.
    """

    item: LegacyItem
    name: str
    kind: str
    input_path: str | None = None
    input_sha256: str | None = None
    image_settings: dict[str, Any] | None = None
    units: dict[tuple[str, int], _Unit] = field(default_factory=dict)
    summaries: dict[int, SummaryRow] = field(default_factory=dict)
    pending: dict[str, list[PlannedUnit]] = field(default_factory=dict)
    venues: dict[str, tuple[str, str]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _latest(log: WorkingLog | None) -> tuple[dict[int, dict[str, Any]], Counter[int]]:
    """Return the last record of every page and the number of records."""
    latest: dict[int, dict[str, Any]] = {}
    counts: Counter[int] = Counter()
    for record in log.results if log is not None else ():
        idx = int(record[PAGE_KEY])
        latest[idx] = record
        counts[idx] += 1
    return latest, counts


def _source_index(kind: str, idx: int, record: Mapping[str, Any]) -> int | str | None:
    page_index = _int(record.get("page_index"))
    if page_index is not None:
        return page_index
    return idx if kind == "pdf" else _text(record.get("image"))


def _transcription_unit(
    record: Mapping[str, Any], idx: int, kind: str, model: str, attempts: int
) -> _Unit:
    raw = record.get("transcription")
    text = "" if raw is None else raw if isinstance(raw, str) else str(raw)
    error = None if record.get("error") is None else str(record["error"])
    unit_model = _text(record.get("model")) or model
    provenance = {
        key: record[key] for key in ("source_file", "image_provenance") if key in record
    }
    return _Unit(
        kind=KIND_TRANSCRIPTION,
        idx=idx,
        status="failed" if error is not None else "done",
        model=unit_model,
        provider=_provider(unit_model, record.get("provider")),
        content=text,
        outcome=transcription_outcome(text, error),
        label=_text(record.get("image")),
        source_index=_source_index(kind, idx, record),
        error=error,
        error_type=_text(record.get("error_type")) if error is not None else None,
        provenance=_plain(provenance) if provenance else None,
        attempts=attempts,
    )


def _summary_unit(
    record: Mapping[str, Any],
    idx: int,
    model: str,
    attempts: int,
    page: _Unit | None,
) -> _Unit:
    transcription_failed = page is not None and page.status == "failed"
    error = None if record.get("error") is None else str(record["error"])
    failed = error is not None and not transcription_failed
    error = error if failed else None
    return _Unit(
        kind=KIND_SUMMARY,
        idx=idx,
        status="failed" if failed else "done",
        model=model,
        provider=_provider(model, record.get("provider")),
        outcome=summary_outcome(
            page.content if page is not None else None, error, transcription_failed
        ),
        label=page.label if page is not None else _text(record.get("image_filename")),
        source_index=page.source_index if page is not None else None,
        error=error,
        error_type=_text(record.get("error_type")) if failed else None,
        attempts=attempts,
    )


def _plan_missing(
    document: LegacyDocument, total: int | None, summarized: bool
) -> None:
    if total is None:
        return
    kinds = (KIND_TRANSCRIPTION, KIND_SUMMARY) if summarized else (KIND_TRANSCRIPTION,)
    for kind in kinds:
        document.pending[kind] = [
            PlannedUnit(idx, source_index=idx if document.kind == "pdf" else None)
            for idx in range(total)
            if (kind, idx) not in document.units
        ]


def _header_of(logs: Sequence[WorkingLog | None]) -> dict[str, Any]:
    for log in logs:
        if log is not None:
            return log.header
    return {}


def legacy_document(item: LegacyItem) -> LegacyDocument | None:
    """Return what the store should hold for *item*; None without log records.

    Raises:
        WorkingLogError: A log lacks the version 2 header.
        OSError: A log cannot be read.
        UnicodeDecodeError: A log is not UTF-8.
    """
    transcription_path, summary_path = item.logs
    transcription_log = read_working_log(transcription_path)
    summary_log = read_working_log(summary_path)
    if transcription_log is None and summary_log is None:
        return None
    header = _header_of((transcription_log, summary_log))
    warnings: list[str] = []
    name = _text(header.get("input_item_name")) or item.base
    if item_base(name, item.root) != item.base:
        warnings.append(
            f"the logs name the item {name!r}, which does not render as "
            f"{item.base!r}; it is imported as {item.base!r}"
        )
        name = item.base
    provenance = _mapping(header.get("file_provenance"))
    config = provenance.get("image_config")
    document = LegacyDocument(
        item=item,
        name=name,
        kind="pdf" if header.get("input_type") == "PDF" else "images",
        input_path=_text(header.get("input_item_path")),
        input_sha256=_text(provenance.get("source_sha256")),
        image_settings=(
            image_settings_record(config) if isinstance(config, Mapping) else None
        ),
        warnings=warnings,
    )
    for log in (transcription_log, summary_log):
        if log is not None and log.dropped_tail:
            warnings.append("a truncated last line of a working log was dropped")
    model = _text(_header_of((transcription_log,)).get("model_name")) or ""
    latest, counts = _latest(transcription_log)
    for idx, record in sorted(latest.items()):
        document.units[(KIND_TRANSCRIPTION, idx)] = _transcription_unit(
            record, idx, document.kind, model, counts[idx]
        )
    summary_model = _text(_header_of((summary_log,)).get("model_name")) or ""
    latest, counts = _latest(summary_log)
    for idx, record in sorted(latest.items()):
        page = document.units.get((KIND_TRANSCRIPTION, idx))
        document.units[(KIND_SUMMARY, idx)] = _summary_unit(
            record, idx, summary_model, counts[idx], page
        )
        document.summaries[idx] = summary_row(idx, record)
    document.venues = {
        KIND_TRANSCRIPTION: (model, _provider(model)),
        KIND_SUMMARY: (summary_model, _provider(summary_model)),
    }
    _plan_missing(document, _int(header.get("total_images")), summary_log is not None)
    return document


# =============================================================================
# Reports
# =============================================================================


@dataclass
class DocumentReport:
    """The import and verification of one document.

    *units* and *failures* describe the log records; *imported*,
    *citations* and *outputs* count what was (or, in a dry run, would be)
    written; *kept* counts units a run of the store wrote, which the import
    leaves as they are. *before* and *after* are the document's resume
    states (``versionless`` while its database predates the store,
    ``unreadable`` for a database the store refuses).
    """

    root: Path
    name: str
    source: Path | None = None
    units: int = 0
    failures: int = 0
    imported: int = 0
    kept: int = 0
    citations: int = 0
    outputs: int = 0
    before: str | None = None
    after: str | None = None
    mismatches: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def changed(self) -> bool:
        """Whether the import writes anything for this document."""
        return bool(self.imported or self.citations or self.outputs)


@dataclass
class DatabaseReport:
    """What migrate did, or would do, with one database.

    *action* is ``convert`` (a dry run's conversion), ``converted``,
    ``create``, ``created``, ``update``, ``updated`` or ``unchanged``.
    """

    path: Path
    action: str
    snapshot: Path | None = None
    documents: int = 0


@dataclass
class MigrationReport:
    """The outcome of :func:`migrate`."""

    input: Path
    dry_run: bool
    databases: list[DatabaseReport] = field(default_factory=list)
    documents: list[DocumentReport] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def mismatches(self) -> int:
        """The number of verification mismatches."""
        return sum(len(document.mismatches) for document in self.documents)

    @property
    def exit_code(self) -> int:
        """1 for errors or mismatches, else 0."""
        return 1 if self.errors or self.mismatches else 0

    def total(self, name: str) -> int:
        """Return the sum of a document counter over all documents."""
        return sum(int(getattr(document, name)) for document in self.documents)

    @property
    def snapshots(self) -> list[Path]:
        """The snapshots written."""
        return [db.snapshot for db in self.databases if db.snapshot is not None]

    def matrix(self) -> dict[tuple[str, str], int]:
        """Count the documents per resume state before and after the import."""
        counts = Counter(
            (document.before or "none", document.after or "not checked")
            for document in self.documents
        )
        return dict(sorted(counts.items()))


# =============================================================================
# Planning
# =============================================================================


@dataclass
class _Existing:
    """What the store already holds for one document."""

    doc: dict[str, Any] | None = None
    units: dict[tuple[str, int], dict[str, Any]] = field(default_factory=dict)
    summaries: dict[int, dict[str, Any]] = field(default_factory=dict)
    citations: int = 0
    outputs: set[str] = field(default_factory=set)


@dataclass
class _StoreState:
    migrate_runs: set[str] = field(default_factory=set)
    documents: dict[str, _Existing] = field(default_factory=dict)


def _existing(connection: sqlite3.Connection, doc: dict[str, Any]) -> _Existing:
    doc_id = doc["doc_id"]
    existing = _Existing(doc=doc)
    for row in _rows(connection, "SELECT * FROM units WHERE doc_id = ?", (doc_id,)):
        existing.units[(row["kind"], row["idx"])] = row
    for row in _rows(connection, "SELECT * FROM summaries WHERE doc_id = ?", (doc_id,)):
        existing.summaries[row["idx"]] = row
    (existing.citations,) = connection.execute(
        "SELECT count(*) FROM citations WHERE doc_id = ?", (doc_id,)
    ).fetchone()
    existing.outputs = {
        str(path)
        for (path,) in connection.execute(
            "SELECT path FROM outputs WHERE doc_id = ?", (doc_id,)
        )
    }
    return existing


def _read_state(connection: sqlite3.Connection) -> _StoreState:
    state = _StoreState()
    state.migrate_runs = {
        str(run_id)
        for (run_id,) in connection.execute(
            "SELECT run_id FROM runs WHERE command IN (?, ?)",
            (MIGRATE_COMMAND, CONVERTED_COMMAND),
        )
    }
    for doc in _rows(connection, "SELECT * FROM documents ORDER BY name"):
        state.documents[str(doc["name"])] = _existing(connection, doc)
    return state


@dataclass
class _Changes:
    """What the import writes for one document."""

    root: Path
    name: str
    report: DocumentReport
    document: LegacyDocument | None = None
    kind: str | None = None
    fields: dict[str, str] = field(default_factory=dict)
    image_settings: dict[str, Any] | None = None
    transcriptions: list[_Unit] = field(default_factory=list)
    summaries: list[tuple[_Unit, SummaryRow]] = field(default_factory=list)
    pending: dict[str, list[PlannedUnit]] = field(default_factory=dict)
    citations: list[CitationRow] = field(default_factory=list)
    outputs: list[tuple[Path, str]] = field(default_factory=list)
    kept: set[tuple[str, int]] = field(default_factory=set)
    create: bool = False
    seen: dict[tuple[str, int], tuple[Any, ...] | None] = field(default_factory=dict)
    """The stored state of each unit the import writes, as the plan read it."""
    seen_citations: int = 0

    @property
    def empty(self) -> bool:
        return not (
            self.create
            or self.fields
            or self.image_settings is not None
            or self.transcriptions
            or self.summaries
            or any(self.pending.values())
            or self.citations
            or self.outputs
        )


def _merged(unit: _Unit, row: Mapping[str, Any] | None) -> _Unit:
    """Return *unit* with the stored usage and model where the logs have none."""
    if row is None:
        return unit
    if unit.usage is None:
        unit.usage = _json_value(row.get("usage"))
    if not unit.model:
        unit.model = str(row.get("model") or "")
    return unit


def _summary_differs(row: SummaryRow, stored: Mapping[str, Any] | None) -> bool:
    if stored is None:
        return True
    return (
        row.page != stored.get("page")
        or _plain(row.page_information) != _json_value(stored.get("page_information"))
        or _plain(row.bullet_points) != _json_value(stored.get("bullet_points"))
        or _plain(row.references) != _json_value(stored.get("page_references"))
    )


def _identity(row: Mapping[str, Any] | None) -> tuple[Any, ...] | None:
    """Return what identifies a unit row's state: status, run, time and stale."""
    if row is None:
        return None
    return (row["status"], row["run_id"], row["updated_at"], row["stale"])


def _unit_changes(
    changes: _Changes, document: LegacyDocument, existing: _Existing, runs: set[str]
) -> None:
    retouched: set[int] = set()
    # Transcriptions first: a page whose transcription changes is retouched,
    # and its summary is then imported again.
    ordered = sorted(
        document.units.items(),
        key=lambda item: (item[0][0] != KIND_TRANSCRIPTION, item[0]),
    )
    for key, unit in ordered:
        row = existing.units.get(key)
        if row is not None and row["status"] != "pending" and row["run_id"] not in runs:
            changes.kept.add(key)
            continue
        merged = _merged(unit, row)
        changes.seen[key] = _identity(row)
        if unit.kind == KIND_TRANSCRIPTION:
            if row is None or merged.differs(row):
                changes.transcriptions.append(merged)
            if row is not None and (
                row["status"],
                row["content"],
            ) != (merged.status, merged.content):
                retouched.add(unit.idx)
            continue
        # A summary depends on its page's transcription: a run that changes
        # it after the plan makes the imported summary outdated.
        page = (KIND_TRANSCRIPTION, unit.idx)
        changes.seen.setdefault(page, _identity(existing.units.get(page)))
        summary = document.summaries[unit.idx]
        if (
            row is None
            or merged.differs(row)
            or _summary_differs(summary, existing.summaries.get(unit.idx))
            or unit.idx in retouched
            or row["stale"]
        ):
            changes.summaries.append((merged, summary))
    for kind, planned in document.pending.items():
        changes.pending[kind] = [
            unit for unit in planned if (kind, unit.idx) not in existing.units
        ]


def _summary_entries(changes: _Changes, existing: _Existing) -> list[dict[str, Any]]:
    """Return the stored summaries as the import leaves them, in page order."""
    entries: dict[int, dict[str, Any]] = {}
    for (kind, idx), row in existing.units.items():
        stored = existing.summaries.get(idx)
        if kind == KIND_SUMMARY and row["status"] != "pending" and stored is not None:
            entries[idx] = {
                "original_input_order_index": idx,
                "image_filename": row["label"],
                "page": stored["page"],
                "page_information": _json_value(stored["page_information"]),
                "bullet_points": _json_value(stored["bullet_points"]),
                "references": _json_value(stored["page_references"]),
                **(
                    {"error": row["error"] or "failed"}
                    if row["status"] == "failed"
                    else {}
                ),
            }
    for unit, summary in changes.summaries:
        entries[unit.idx] = {
            "original_input_order_index": unit.idx,
            "image_filename": unit.label,
            "page": summary.page,
            "page_information": _plain(summary.page_information),
            "bullet_points": _plain(summary.bullet_points),
            "references": _plain(summary.references),
            **({"error": unit.error or "failed"} if unit.status == "failed" else {}),
        }
    return [entries[idx] for idx in sorted(entries)]


def _resume_state(root: Path, name: str, document: LegacyDocument | None) -> str:
    """Return the resume state of *name* as a run would classify it."""
    outputs = [path for path, _fmt in item_files(name, root)]
    legacy = ItemPaths.for_item(name, root).legacy_logs()
    try:
        state = classify(root, name, ResumeRequest(outputs=outputs, legacy=legacy))
    except StoreError:
        return "unreadable"
    except InputMismatchError:
        return "mismatch"
    return state.state.value


def _plan_document(
    root: Path,
    name: str,
    document: LegacyDocument | None,
    existing: _Existing,
    state: _StoreState,
    cache: OpenAlexCache,
) -> _Changes:
    report = DocumentReport(root, name, document.item.folder if document else None)
    changes = _Changes(root, name, report, document)
    doc = existing.doc
    if document is not None:
        report.warnings.extend(document.warnings)
        changes.create = doc is None
        changes.kind = document.kind
        for column, value in (
            ("input_path", document.input_path),
            ("input_sha256", document.input_sha256),
        ):
            if value is not None and (doc is None or doc.get(column) is None):
                changes.fields[column] = value
        if document.image_settings is not None and (
            doc is None or doc.get("image_settings") is None
        ):
            changes.image_settings = document.image_settings
        _unit_changes(changes, document, existing, state.migrate_runs)
    changes.outputs = [
        (path, fmt)
        for path, fmt in item_files(name, root)
        if _output_key(path, root) not in existing.outputs
    ]
    changes.seen_citations = existing.citations
    owned = doc is None or doc.get("first_run_id") in state.migrate_runs
    if owned and not existing.citations:
        kind = changes.kind or (doc or {}).get("kind")
        label = "pdf" if kind == "pdf" else "image"
        entries = _summary_entries(changes, existing)
        if entries:
            changes.citations = offline_citations(entries, label, cache)
    report.kept = len(changes.kept)
    report.imported = (
        len(changes.transcriptions)
        + len(changes.summaries)
        + sum(len(planned) for planned in changes.pending.values())
    )
    report.citations = len(changes.citations)
    report.outputs = len(changes.outputs)
    if document is not None:
        report.units = len(document.units)
        report.failures = sum(u.status == "failed" for u in document.units.values())
    return changes


@dataclass
class _Root:
    root: Path
    state: SchemaState
    documents: list[LegacyDocument] = field(default_factory=list)
    before: dict[str, str] = field(default_factory=dict)


def _plan_root(entry: _Root, cache: OpenAlexCache) -> list[_Changes]:
    database = database_path(entry.root)
    if entry.state.exists and not entry.state.empty and not entry.state.versionless:
        state = read(database, _read_state)
    else:
        state = _StoreState()
    planned = []
    logged = {document.name for document in entry.documents}
    for document in entry.documents:
        existing = state.documents.get(document.name, _Existing())
        planned.append(
            _plan_document(entry.root, document.name, document, existing, state, cache)
        )
    for name, existing in state.documents.items():
        doc = existing.doc or {}
        if name not in logged and doc.get("first_run_id") in state.migrate_runs:
            planned.append(
                _plan_document(entry.root, name, None, existing, state, cache)
            )
    for changes in planned:
        changes.report.before = entry.before.get(changes.name) or _resume_state(
            entry.root, changes.name, changes.document
        )
    return planned


# =============================================================================
# Writing
# =============================================================================


class _ChangedError(Exception):
    """The document changed between the plan and the leased write."""


def _revalidate(connection: sqlite3.Connection, changes: _Changes, doc_id: int) -> None:
    """Raise :class:`_ChangedError` when the planned rows changed since the plan."""
    current = {
        (row["kind"], row["idx"]): _identity(row)
        for row in _rows(
            connection,
            "SELECT kind, idx, status, run_id, updated_at, stale FROM units "
            "WHERE doc_id = ?",
            (doc_id,),
        )
    }
    changed = sorted(
        key for key, seen in changes.seen.items() if current.get(key) != seen
    )
    (citations,) = connection.execute(
        "SELECT count(*) FROM citations WHERE doc_id = ?", (doc_id,)
    ).fetchone()
    if changed or citations != changes.seen_citations:
        units = ", ".join(f"{kind} {idx}" for kind, idx in changed) or "citations"
        raise _ChangedError(
            f"'{changes.name}' changed while migrate planned it ({units}); it was "
            "not imported. Run migrate again."
        )


def _write_document(
    connection: sqlite3.Connection, changes: _Changes, doc_id: int, run_id: str
) -> None:
    _revalidate(connection, changes, doc_id)
    now = _now()
    connection.execute(
        "UPDATE documents SET input_path = COALESCE(input_path, ?), "
        "input_sha256 = COALESCE(input_sha256, ?), kind = COALESCE(kind, ?) "
        "WHERE doc_id = ?",
        (
            changes.fields.get("input_path"),
            changes.fields.get("input_sha256"),
            changes.kind,
            doc_id,
        ),
    )
    if changes.image_settings is not None:
        connection.execute(
            "UPDATE documents SET image_settings = ? WHERE doc_id = ? "
            "AND image_settings IS NULL",
            (_json(changes.image_settings), doc_id),
        )
    for unit in changes.transcriptions:
        connection.execute(_UPSERT_UNIT, unit.params(doc_id, run_id, now))
    for unit, row in changes.summaries:
        connection.execute(
            _UPSERT_SUMMARY,
            (
                doc_id,
                unit.idx,
                row.page,
                _json(row.page_information),
                _json(row.bullet_points),
                _json(row.references),
            ),
        )
        connection.execute(_UPSERT_UNIT, unit.params(doc_id, run_id, now))
    document = changes.document
    for kind, planned in changes.pending.items():
        if document is None or not planned:
            continue
        model, provider = document.venues[kind]
        connection.executemany(
            _PLAN_UNIT,
            [
                (doc_id, kind, unit.idx, unit.source_index, unit.label, TOOL)
                + (model, provider, run_id, now)
                for unit in planned
            ],
        )
    if changes.citations:
        connection.execute("DELETE FROM citations WHERE doc_id = ?", (doc_id,))
        connection.executemany(
            "INSERT INTO citations (doc_id, position, text, pages, partial, doi, "
            "url, metadata) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            [
                (
                    doc_id,
                    position,
                    citation.text,
                    citation.pages,
                    int(citation.partial),
                    citation.doi,
                    citation.url,
                    _json(citation.metadata),
                )
                for position, citation in enumerate(changes.citations, start=1)
            ],
        )
    refresh_documents(connection, doc_id)
    if changes.transcriptions or changes.summaries or any(changes.pending.values()):
        connection.execute(
            "UPDATE documents SET last_run_id = ?, updated_at = ? WHERE doc_id = ?",
            (run_id, now, doc_id),
        )
    for path, fmt in changes.outputs:
        _register_output(connection, doc_id, path, fmt, changes.root, run_id)


async def _apply(store: Store, changes: _Changes, run_id: str) -> None:
    """Write one document's changes under its lease.

    The plan is checked again inside the leased write: a document a run
    changed since the plan is left as it is and reported as a mismatch. Only
    the document's name is ensured before the lease; its input fields and
    kind are filled, where missing, inside the leased write.
    """
    doc_id = await store.ensure_document(changes.name)
    lease = await store.acquire(doc_id)
    if lease is None:
        message = f"'{changes.name}' is leased by another run; it was not imported."
        await store.add_note(message, doc_id=doc_id)
        changes.report.mismatches.append(message)
        return
    try:
        await store.write(
            lease,
            lambda connection: _write_document(connection, changes, doc_id, run_id),
        )
    except _ChangedError as exc:
        await store.add_note(str(exc), doc_id=doc_id)
        changes.report.mismatches.append(str(exc))
    finally:
        await store.release(lease)


async def _write_root(
    entry: _Root, planned: Sequence[_Changes], spec: Mapping[str, Any]
) -> None:
    async with open_store(entry.root) as store:
        run_id = await store.begin_run(MIGRATE_COMMAND, spec=spec)
        code = 1
        try:
            for changes in planned:
                if not changes.empty:
                    await _apply(store, changes, run_id)
            code = 0
        finally:
            await store.end_run(code)


# =============================================================================
# Verification
# =============================================================================


def _verify_bodies(
    connection: sqlite3.Connection,
    doc_id: int,
    outputs: Sequence[tuple[Path, str]],
    report: DocumentReport,
) -> None:
    for path, fmt in outputs:
        problem = body_problem(connection, doc_id, path, fmt)
        if problem is not None:
            report.mismatches.append(problem)


def _verify_units(
    changes: _Changes,
    units: Mapping[tuple[str, int], Mapping[str, Any]],
    report: DocumentReport,
) -> None:
    document = changes.document
    if document is None:
        return
    for kind in (KIND_TRANSCRIPTION, KIND_SUMMARY):
        expected = [
            (key, unit)
            for key, unit in sorted(document.units.items())
            if key[0] == kind and key not in changes.kept
        ]
        matching = failed = recorded_failures = 0
        for key, unit in expected:
            row = units.get(key)
            recorded_failures += unit.status == "failed"
            if row is None:
                report.mismatches.append(f"{kind} unit {key[1]} is not in the store")
                continue
            failed += row["status"] == "failed"
            same = row["status"] == unit.status
            if kind == KIND_TRANSCRIPTION:
                same = same and row["content"] == unit.content
            matching += same
        if matching != len(expected):
            report.mismatches.append(
                f"{matching} of {len(expected)} {kind} units match their log records"
            )
        if failed != recorded_failures:
            report.mismatches.append(
                f"{failed} failed {kind} units stored, {recorded_failures} recorded"
            )
        for planned in document.pending.get(kind, ()):
            if (kind, planned.idx) not in units:
                report.mismatches.append(
                    f"pending {kind} unit {planned.idx} is missing"
                )


def _verify_document(connection: sqlite3.Connection, changes: _Changes) -> None:
    report = changes.report
    docs = _rows(connection, "SELECT * FROM documents WHERE name = ?", (changes.name,))
    if not docs:
        report.mismatches.append(f"'{changes.name}' is not in the store")
        return
    doc = docs[0]
    doc_id = int(doc["doc_id"])
    units = {
        (row["kind"], row["idx"]): row
        for row in _rows(connection, "SELECT * FROM units WHERE doc_id = ?", (doc_id,))
    }
    _verify_units(changes, units, report)
    document = changes.document
    if (
        document is not None
        and document.image_settings is not None
        and doc["image_settings"] is None
    ):
        report.mismatches.append("the image-settings record is missing")
    registered = {
        str(path): str(fmt)
        for path, fmt in connection.execute(
            "SELECT path, format FROM outputs WHERE doc_id = ?", (doc_id,)
        )
    }
    files = item_files(changes.name, changes.root)
    for path, _fmt in files:
        if _output_key(path, changes.root) not in registered:
            report.mismatches.append(f"{path.name} is not registered")
    _verify_bodies(connection, doc_id, files, report)


def _verify_root(entry: _Root, planned: Sequence[_Changes]) -> None:
    database = database_path(entry.root)
    if not database.is_file():
        return

    def verify(connection: sqlite3.Connection) -> None:
        for changes in planned:
            _verify_document(connection, changes)

    read(database, verify)
    for changes in planned:
        changes.report.after = _resume_state(entry.root, changes.name, changes.document)


# =============================================================================
# migrate
# =============================================================================


def _versionless_names(database: Path) -> list[str]:
    return [
        str(name)
        for (name,) in read(
            database,
            lambda connection: connection.execute(
                "SELECT name FROM documents"
            ).fetchall(),
        )
    ]


def _roots(path: Path, report: MigrationReport) -> dict[Path, _Root]:
    """Find the items and databases of *path*; check every database first."""
    path = Path(path).resolve()
    roots: dict[Path, _Root] = {}

    def add(root: Path) -> _Root:
        key = root.resolve()
        if key not in roots:
            roots[key] = _Root(key, _check_database(database_path(key)))
        return roots[key]

    if path.is_file() and path.suffix.lower() == ".sqlite":
        add(path.parent)
    elif path.is_dir():
        for database in sorted(path.rglob(DATABASE_NAME)):
            if database.is_file():
                add(database.parent)
    for item in find_items(path):
        entry = add(item.root)
        try:
            document = legacy_document(item)
        except (WorkingLogError, OSError, UnicodeDecodeError) as exc:
            report.errors.append(f"{item.folder}: {exc}")
            continue
        if document is not None:
            entry.documents.append(document)
    for entry in roots.values():
        if entry.state.versionless:
            names = {*_versionless_names(database_path(entry.root))}
            names |= {document.name for document in entry.documents}
            entry.before = dict.fromkeys(sorted(names), "versionless")
    return roots


async def migrate(
    path: Path, *, dry_run: bool = False, state_dir: Path | None = None
) -> MigrationReport:
    """Import the version 1.0 outputs at or below *path* into their stores.

    With *dry_run* nothing is written: the report lists what the import
    would write. *state_dir* holds the OpenAlex cache the citations take
    their metadata from (None: the default state directory); no request is
    made.

    Raises:
        StoreError: A database is newer, unreadable or of an unknown
            versionless layout; nothing was written.
        FileNotFoundError: *path* does not exist.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Input path does not exist: {path}")
    report = MigrationReport(path.resolve(), dry_run)
    roots = _roots(path, report)
    cache = OpenAlexCache(state_dir)
    spec = {
        "input": {"value": str(report.input), "source": "flag"},
        "dry_run": {"value": dry_run, "source": "flag"},
    }
    for entry in roots.values():
        database = database_path(entry.root)
        db_report = DatabaseReport(database, "unchanged")
        report.databases.append(db_report)
        if entry.state.versionless:
            db_report.action = "convert"
            if not dry_run:
                saved = snapshot_path(database)
                try:
                    convert_versionless_ae(database, saved)
                except MigrationError as exc:
                    report.errors.append(str(exc))
                    continue
                db_report.action = "converted"
                db_report.snapshot = saved
                entry.state = _check_database(database)
        planned = _plan_root(entry, cache)
        report.documents.extend(changes.report for changes in planned)
        db_report.documents = len(planned)
        pending = [changes for changes in planned if not changes.empty]
        if pending and db_report.action == "unchanged":
            db_report.action = "update" if entry.state.exists else "create"
        if dry_run:
            continue
        if pending:
            if entry.state.exists and db_report.snapshot is None:
                db_report.snapshot = snapshot_path(database)
                snapshot(database, db_report.snapshot)
            await _write_root(entry, planned, spec)
            if db_report.action in ("update", "create"):
                db_report.action += "d"
        _verify_root(entry, planned)
    return report
