"""Round-3 rendering regression tests.

One focused regression per verified fix in the rendering layer and its
supporting constants:

1.  Citation ed/eds/trans stripping no longer fires inside words / compounds.
2.  A content bullet mentioning "error" is no longer dropped.
3.  ``is_blank_transcription`` is length-gated (long prose is not "blank").
4.  Failed-page error placeholders stay visible in the rendered outputs.
5.  Non-string bullet entries are filtered instead of crashing the writers.
6.  Internal newlines in bullets/citations are collapsed at render time.
7.  ``_extract_volume`` no longer reads an author initial as a tome/volume.
8.  The DOCX writer saves atomically (temp file + os.replace).
9.  The "total pages" metadata reports the pre-filter source page count.
10. The text writer coerces non-string transcriptions and cleans up its temp.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
from unittest.mock import patch

from docx import Document

from config.constants import is_blank_transcription
from rendering.citations import Citation, CitationManager
from rendering.markdown import create_markdown_summary
from rendering.summary import (
    _is_meaningful_summary,
    collapse_internal_newlines,
    prepare_summary_data,
)
from rendering.text import write_transcription_to_text


def _content_page(
    page_num: int, bullets: object, page_types: list[str] | None = None
) -> dict[str, Any]:
    """Build a minimal flat summary result for a content page."""
    return {
        "page": page_num,
        "page_information": {
            "page_number_integer": page_num,
            "is_two_page_spread": False,
            "page_number_integer_end": None,
            "page_number_type": "arabic",
            "page_types": page_types or ["content"],
        },
        "bullet_points": bullets,
        "references": None,
    }


# ---------------------------------------------------------------------------
# Item 1: ed/eds/trans stripping must not fire inside words or compounds
# ---------------------------------------------------------------------------
class TestEditorMarkerStripping:
    def test_education_keeps_its_ed(self) -> None:
        """'Education' must not lose its leading 'ed' to the editor-marker regex."""
        citation = Citation(raw_text="Smith, J. 1990. Education Reform.")
        assert "education" in citation.comparison_text

    def test_trans_atlantic_pair_not_byte_identical(self) -> None:
        """'Trans-Atlantic' and 'Atlantic' no longer collapse to one citation."""
        trans = Citation(raw_text="Smith, J. (1990). Trans-Atlantic Trade Networks.")
        atlantic = Citation(raw_text="Smith, J. (1990). Atlantic Trade Networks.")
        # The hyphenated compound keeps its "trans", so the comparison texts and
        # thus the dedup keys differ (a bare \b would still strip it).
        assert "trans" in trans.comparison_text
        assert trans.comparison_text != atlantic.comparison_text
        assert trans.normalized_key != atlantic.normalized_key

        manager = CitationManager()
        manager.add_citations([trans.raw_text], 1)
        manager.add_citations([atlantic.raw_text], 1)
        # Previously these merged deterministically into a single citation.
        assert len(manager.citations) == 2

    def test_distinct_titles_survive_consolidate(self) -> None:
        """Genuinely distinct same-author/year works are not fuzzy-merged."""
        manager = CitationManager()
        manager.add_citations(
            [
                "Smith, J. (1990). Trans-Atlantic Trade Networks in the Early "
                "Modern Atlantic World. Oxford University Press."
            ],
            1,
        )
        manager.add_citations(
            [
                "Smith, J. (1990). Atlantic Migration and Labor Systems, "
                "1500-1800. Cambridge University Press."
            ],
            2,
        )
        manager.consolidate()
        assert len(manager.citations) == 2

    def test_standalone_editor_marker_still_stripped(self) -> None:
        """A real '(ed.)' marker is still removed from the comparison text."""
        citation = Citation(raw_text="Jones, A. (ed.). 1985. Collected Essays.")
        assert "ed" not in citation.comparison_text.split()


# ---------------------------------------------------------------------------
# Item 2: a content bullet that merely mentions "error" is kept
# ---------------------------------------------------------------------------
class TestMeasurementErrorBulletKept:
    def test_error_word_in_content_bullet_is_meaningful(self) -> None:
        summary = _content_page(
            1, ["Discusses measurement error in historical wage series."]
        )
        assert _is_meaningful_summary(summary) is True

    def test_bracketed_empty_page_marker_still_dropped(self) -> None:
        """A single bracketed blank sentinel on a content page is still dropped."""
        summary = _content_page(1, ["[empty page]"])
        assert _is_meaningful_summary(summary) is False


# ---------------------------------------------------------------------------
# Item 3: is_blank_transcription is length-gated
# ---------------------------------------------------------------------------
class TestBlankTranscriptionLengthGate:
    def test_long_prose_mentioning_empty_page_is_not_blank(self) -> None:
        prose = (
            "The author reflects at length on the metaphor of the empty page "
            "as a site of possibility. " * 40
        )
        assert len(prose) > 2000
        assert is_blank_transcription(prose) is False

    def test_short_bracketed_sentinel_is_blank(self) -> None:
        sentinel = "[<img>: no transcribable text — blank scan]"
        assert is_blank_transcription(sentinel) is True


# ---------------------------------------------------------------------------
# Item 4: failed-page error placeholders stay visible
# ---------------------------------------------------------------------------
class TestErrorPlaceholderVisible:
    _ERROR_RESULT = {
        "page": 3,
        "page_information": {
            "page_number_integer": 3,
            "is_two_page_spread": False,
            "page_number_integer_end": None,
            "page_number_type": "arabic",
            "page_types": ["other"],
        },
        "bullet_points": ["[Error generating summary: request timed out]"],
        "references": None,
        "error": "request timed out",
    }

    def test_error_page_is_meaningful_and_rendered(self) -> None:
        assert _is_meaningful_summary(self._ERROR_RESULT) is True
        manager = CitationManager()
        data = prepare_summary_data([self._ERROR_RESULT], manager)
        # The failed page produces a visible render item with its placeholder.
        assert len(data.page_render_items) == 1
        item = data.page_render_items[0]
        assert item.bullet_points == ["[Error generating summary: request timed out]"]
        assert "Page 3" in item.heading_text

    def test_error_page_visible_in_markdown(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.md"
        with patch("rendering.markdown.enrich_if_enabled"):
            create_markdown_summary([self._ERROR_RESULT], output_path, "Doc")
        content = output_path.read_text(encoding="utf-8")
        assert "[Error generating summary: request timed out]" in content


# ---------------------------------------------------------------------------
# Item 5: non-string bullet entries are filtered, not crashed on
# ---------------------------------------------------------------------------
class TestNonStringBulletsFiltered:
    _RESULT = _content_page(1, [None, 123, "real bullet"])

    def test_meaningful_and_single_string_bullet(self) -> None:
        assert _is_meaningful_summary(self._RESULT) is True
        manager = CitationManager()
        data = prepare_summary_data([self._RESULT], manager)
        assert data.page_render_items[0].bullet_points == ["real bullet"]

    def test_markdown_renders_only_string_bullet(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.md"
        with patch("rendering.markdown.enrich_if_enabled"):
            create_markdown_summary([self._RESULT], output_path, "Doc")
        content = output_path.read_text(encoding="utf-8")
        assert "- real bullet" in content
        assert "123" not in content

    def test_docx_renders_only_string_bullet(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.docx"
        from rendering.docx import create_docx_summary

        with patch("rendering.docx.enrich_if_enabled"):
            create_docx_summary([self._RESULT], output_path, "Doc")
        assert output_path.exists()
        doc = Document(str(output_path))
        texts = [p.text for p in doc.paragraphs]
        assert any("real bullet" in t for t in texts)


# ---------------------------------------------------------------------------
# Item 6: internal newlines collapsed at render time
# ---------------------------------------------------------------------------
class TestNewlineCollapse:
    def test_helper_collapses_internal_newline(self) -> None:
        assert collapse_internal_newlines("Line one\nLine two") == "Line one Line two"
        assert collapse_internal_newlines("a  \n  b") == "a b"

    def test_markdown_bullet_newline_collapsed(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.md"
        result = _content_page(1, ["First part\nSecond part"])
        with patch("rendering.markdown.enrich_if_enabled"):
            create_markdown_summary([result], output_path, "Doc")
        content = output_path.read_text(encoding="utf-8")
        assert "- First part Second part" in content

    def test_markdown_citation_newline_collapsed(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.md"
        result = {
            "page": 1,
            "page_information": {
                "page_number_integer": 1,
                "is_two_page_spread": False,
                "page_number_integer_end": None,
                "page_number_type": "arabic",
                "page_types": ["content"],
            },
            "bullet_points": ["ok"],
            "references": [
                {
                    "citation": "Smith, J. (2020). *A Long\nBroken Title*. Press.",
                    "is_partial": False,
                }
            ],
        }
        with patch("rendering.markdown.enrich_if_enabled"):
            create_markdown_summary([result], output_path, "Doc")
        content = output_path.read_text(encoding="utf-8")
        assert "A Long Broken Title" in content


# ---------------------------------------------------------------------------
# Item 7: an author initial is not read as a volume
# ---------------------------------------------------------------------------
class TestVolumeInitialFalsePositive:
    def test_initial_before_three_digit_number_not_a_volume(self) -> None:
        assert Citation(raw_text="Smith, T. 190. Title.").volume is None

    def test_initial_before_year_not_a_volume(self) -> None:
        assert Citation(raw_text="Smith, T. 1990. The Wealth. CUP.").volume is None

    def test_real_lowercase_tome_still_parses(self) -> None:
        assert Citation(raw_text="Dupont. *Oeuvres*, t. II. Paris.").volume == 2

    def test_english_volume_still_parses(self) -> None:
        assert Citation(raw_text="Smith (2020). *W*, vol. 3. CUP.").volume == 3


# ---------------------------------------------------------------------------
# Item 8: atomic DOCX save leaves no orphan temp file
# ---------------------------------------------------------------------------
class TestAtomicDocxSave:
    def test_real_save_produces_file_without_temp(self, tmp_path: Path) -> None:
        from rendering.docx import create_docx_summary

        output_path = tmp_path / "summary.docx"
        result = _content_page(1, ["A bullet"])
        with patch("rendering.docx.enrich_if_enabled"):
            create_docx_summary([result], output_path, "Doc")
        assert output_path.exists()
        assert output_path.stat().st_size > 0
        assert not output_path.with_name(output_path.name + ".tmp").exists()


# ---------------------------------------------------------------------------
# Item 9: total pages reports the pre-filter source count
# ---------------------------------------------------------------------------
class TestTotalPagesPreFilter:
    def test_total_counts_filtered_blank_page(self, tmp_path: Path) -> None:
        content = _content_page(1, ["Real content bullet"])
        blank = {
            "page": 2,
            "page_information": {
                "page_number_integer": 2,
                "is_two_page_spread": False,
                "page_number_integer_end": None,
                "page_number_type": "arabic",
                "page_types": ["blank"],
            },
            "bullet_points": None,
            "references": None,
        }
        manager = CitationManager()
        data = prepare_summary_data([content, blank], manager)
        assert len(data.filtered_results) == 1
        assert data.source_page_count == 2

        output_path = tmp_path / "summary.md"
        with patch("rendering.markdown.enrich_if_enabled"):
            create_markdown_summary([content, blank], output_path, "Doc")
        text = output_path.read_text(encoding="utf-8")
        assert "2 total pages*" in text


# ---------------------------------------------------------------------------
# Item 10: text writer coerces non-string transcriptions, no orphan temp
# ---------------------------------------------------------------------------
class TestTextWriterCoercion:
    def test_none_transcription_coerced_no_orphan_temp(self, tmp_path: Path) -> None:
        output_path = tmp_path / "out.txt"
        results: list[dict[str, Any]] = [
            {"transcription": None},
            {"transcription": "real text"},
        ]
        ok = write_transcription_to_text(
            results,
            output_path=output_path,
            document_name="Doc",
            item_type="pdf",
            total_elapsed_time=1.0,
            source_path=tmp_path / "src.pdf",
        )
        assert ok is True
        content = output_path.read_text(encoding="utf-8")
        assert "real text" in content
        assert not output_path.with_name(output_path.name + ".tmp").exists()
