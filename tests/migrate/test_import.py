"""migrate imports the working logs and files of version 1.0 items.

The items are built in ``tmp_path`` as version 1.0 left them: the two
working logs, the files with their old header, transcriptions without the
page-break markers and without a final newline. After the import the store
holds every page, the citations and the registered files, and ``render``
replaces the files in place.
"""

from __future__ import annotations

import asyncio
import io
import json
import sqlite3
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import parse_xml

import autoexcerpter.migrate as migrate_module
from autoexcerpter.common.sqlite import read_rows
from autoexcerpter.common.store import file_sha256
from autoexcerpter.migrate import MigrationReport, body_matches, docx_text, migrate
from autoexcerpter.output.render import render_database
from autoexcerpter.output.store import DATABASE_NAME, read_only_connection
from tests.migrate.conftest import METADATA
from tests.migrate.legacy import (
    CITATION,
    IMAGE_CONFIG,
    MODEL,
    SUMMARY_MODEL,
    Item,
    Page,
    make_item,
    write_files,
    write_logs,
)

MATH = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _pages() -> list[Page]:
    return [
        Page("First page.", references=(CITATION,)),
        Page("An unnumbered page.", numbered=False),
        Page("Third page.", references=(CITATION,)),
    ]


@pytest.fixture
def item(tmp_path: Path, state_dir: Path) -> Item:
    """A complete three-page item: logs and files as version 1.0 left them."""
    legacy = make_item(tmp_path / "out", "doc", _pages(), state_dir=state_dir)
    legacy.retries = {2: 1}
    write_logs(legacy)
    write_files(legacy, markers=False, final_newline=False)
    return legacy


def _migrate(root: Path, state_dir: Path, **options: Any) -> MigrationReport:
    return asyncio.run(migrate(root, state_dir=state_dir, **options))


def _view(root: Path, view: str) -> list[dict[str, Any]]:
    with read_only_connection(root) as connection:
        cursor = connection.execute(f"SELECT * FROM {view}")
        names = [column[0] for column in cursor.description]
        return [dict(zip(names, row, strict=True)) for row in cursor.fetchall()]


def _dump(root: Path) -> dict[str, list[dict[str, Any]]]:
    tables = ("documents", "units", "summaries", "citations", "outputs", "runs")
    return {table: read_rows(root / DATABASE_NAME, table) for table in tables}


def test_a_complete_item_is_imported_with_its_units_citations_and_files(
    item: Item, state_dir: Path
) -> None:
    report = _migrate(item.root, state_dir)

    assert report.exit_code == 0
    (document,) = report.documents
    assert document.warnings == []
    assert (document.imported, document.citations, document.outputs) == (6, 1, 3)
    assert (document.before, document.after) == ("legacy", "complete")
    assert report.databases[0].action == "created"
    pages = _view(item.root, "page_text")
    assert [row["status"] for row in pages] == ["done"] * 3
    assert [row["model"] for row in pages] == [MODEL] * 3
    assert pages[1]["text"] == "An unnumbered page."
    summaries = _view(item.root, "page_summaries")
    assert [row["model"] for row in summaries] == [SUMMARY_MODEL] * 3
    assert json.loads(summaries[0]["bullet_points"]) == ["A point of page 1."]
    (citation,) = _view(item.root, "document_citations")
    assert citation["text"] == CITATION
    assert citation["doi"] == METADATA["doi"]
    assert citation["pages"] == "pp. 1, 3"
    outputs = {row["path"]: row["sha256"] for row in _view(item.root, "outputs")}
    paths = item.paths
    for path in (paths.transcription, paths.summary_md, paths.summary_docx):
        assert outputs[path.name] == file_sha256(path)


def test_the_document_records_the_log_header(item: Item, state_dir: Path) -> None:
    _migrate(item.root, state_dir)

    (row,) = _view(item.root, "documents")
    assert json.loads(row["image_settings"]) == IMAGE_CONFIG
    assert row["input_sha256"] == "ab" * 32
    assert row["input_path"] == str(item.root / "doc.pdf")
    units = {(u["kind"], u["idx"]): u for u in _view(item.root, "units")}
    assert units[("transcription", 2)]["attempts"] == 2
    assert json.loads(units[("transcription", 0)]["provenance"])["image_provenance"]


def test_a_second_run_imports_nothing_and_changes_no_row(
    item: Item, state_dir: Path
) -> None:
    _migrate(item.root, state_dir)
    before = _dump(item.root)

    report = _migrate(item.root, state_dir)

    assert report.exit_code == 0
    (document,) = report.documents
    assert (document.imported, document.citations, document.outputs) == (0, 0, 0)
    assert report.databases[0].action == "unchanged"
    assert report.snapshots == []
    assert _dump(item.root) == before


def test_render_replaces_the_registered_files_in_place(
    item: Item, state_dir: Path
) -> None:
    _migrate(item.root, state_dir)

    (result,) = asyncio.run(render_database(item.root / DATABASE_NAME))

    paths = item.paths
    assert result.kept == []
    assert set(result.written) == {
        paths.transcription,
        paths.summary_md,
        paths.summary_docx,
    }
    text = paths.transcription.read_text(encoding="utf-8")
    assert "Total processing time" not in text
    assert '<page_break pdf="2"/>\nAn unnumbered page.' in text
    assert text.endswith("Third page.\n")


def _hashes(root: Path) -> dict[str, str | None]:
    return {row["path"]: row["sha256"] for row in _view(root, "outputs")}


def test_a_changed_body_fails_and_is_not_taken_as_the_tools_file(
    item: Item, state_dir: Path
) -> None:
    transcription = item.paths.transcription
    transcription.write_text(
        transcription.read_text(encoding="utf-8").replace("First", "Edited"),
        encoding="utf-8",
    )
    edited = transcription.read_bytes()

    report = _migrate(item.root, state_dir)

    assert report.exit_code == 1
    (document,) = report.documents
    assert len(document.mismatches) == 1
    assert transcription.name in document.mismatches[0]
    assert _hashes(item.root)[transcription.name] is None
    (result,) = asyncio.run(render_database(item.root / DATABASE_NAME))
    assert result.kept == [transcription]
    assert transcription.read_bytes() == edited
    assert _migrate(item.root, state_dir).exit_code == 1


@pytest.mark.parametrize("edit", ["paragraph", "table", "unreadable"])
def test_a_word_summary_is_compared_by_its_text(
    item: Item, state_dir: Path, edit: str
) -> None:
    docx = item.paths.summary_docx
    if edit == "unreadable":
        docx.write_bytes(b"PK\x03\x04 not a Word file")
    else:
        document = Document(str(docx))
        if edit == "table":
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Added by hand."
        else:
            document.add_paragraph("A paragraph added by hand.")
        document.save(str(docx))

    report = _migrate(item.root, state_dir)

    assert report.exit_code == 1
    (mismatch,) = report.documents[0].mismatches
    assert docx.name in mismatch
    assert _hashes(item.root)[docx.name] is None


def _with_equation(path: Path, symbol: str) -> bytes:
    document = Document(str(path))
    paragraph = document.add_paragraph("Equation:")
    paragraph._p.append(
        parse_xml(
            f'<m:oMath xmlns:m="{MATH}"><m:sSup><m:e><m:r><m:t>{symbol}</m:t>'
            "</m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>"
        )
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_word_text_keeps_the_equations(item: Item) -> None:
    docx = item.paths.summary_docx

    x, y = _with_equation(docx, "x"), _with_equation(docx, "y")

    assert docx_text(x) != docx_text(y)
    assert docx_text(x) != docx_text(docx.read_bytes())


def test_word_text_leaves_the_metadata_line_out(item: Item) -> None:
    lines = docx_text(item.paths.summary_docx.read_bytes())

    assert lines[0] == "doc"
    assert not any(line.startswith("Processed ") for line in lines)
    assert "A point of page 1." in lines


def test_a_unit_a_run_stored_after_the_plan_is_left_and_reported(
    tmp_path: Path, state_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    legacy = make_item(
        tmp_path / "out", "doc", _pages()[:2], state_dir=state_dir, total=3
    )
    write_logs(legacy)
    _migrate(legacy.root, state_dir)
    legacy.pages = _pages()
    write_logs(legacy)
    apply = migrate_module._apply

    async def run_first(store: Any, changes: Any, run_id: str) -> None:
        connection = sqlite3.connect(legacy.root / DATABASE_NAME)
        with connection:
            connection.execute(
                "INSERT INTO runs (run_id, tool, tool_version, command, started_at) "
                "VALUES ('a-run', 'autoexcerpter', '2.0.0', 'run', '2026')"
            )
            connection.execute(
                "UPDATE units SET status = 'done', content = 'From a run.', "
                "run_id = 'a-run', updated_at = '2026' WHERE idx = 2"
            )
        connection.close()
        await apply(store, changes, run_id)

    monkeypatch.setattr(migrate_module, "_apply", run_first)

    report = _migrate(legacy.root, state_dir)

    assert report.exit_code == 1
    assert any(
        "changed while migrate planned it" in m for m in report.documents[0].mismatches
    )
    units = {(u["kind"], u["idx"]): u for u in _view(legacy.root, "units")}
    assert units[("transcription", 2)]["content"] == "From a run."
    assert units[("transcription", 2)]["run_id"] == "a-run"


def _before_apply(
    monkeypatch: pytest.MonkeyPatch, root: Path, *statements: str
) -> None:
    """Run *statements* as another run would, between the plan and the write."""
    apply = migrate_module._apply

    async def run_first(store: Any, changes: Any, run_id: str) -> None:
        connection = sqlite3.connect(root / DATABASE_NAME)
        with connection:
            connection.execute(
                "INSERT OR IGNORE INTO runs (run_id, tool, tool_version, command, "
                "started_at) VALUES ('a-run', 'autoexcerpter', '2.0.0', 'run', '2026')"
            )
            for statement in statements:
                connection.execute(statement)
        connection.close()
        await apply(store, changes, run_id)

    monkeypatch.setattr(migrate_module, "_apply", run_first)


def _changed(report: MigrationReport) -> bool:
    return any(
        "changed while migrate planned it" in m for m in report.documents[0].mismatches
    )


def test_an_input_hash_a_run_recorded_after_the_plan_is_kept(
    item: Item, state_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _migrate(item.root, state_dir)
    connection = sqlite3.connect(item.root / DATABASE_NAME)
    with connection:
        connection.execute("UPDATE documents SET input_sha256 = NULL")
    connection.close()
    _before_apply(monkeypatch, item.root, "UPDATE documents SET input_sha256 = 'newer'")

    _migrate(item.root, state_dir)

    (document,) = _view(item.root, "documents")
    assert document["input_sha256"] == "newer"


def test_a_summary_marked_stale_after_the_plan_is_left_and_reported(
    tmp_path: Path, state_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    legacy = make_item(
        tmp_path / "out", "doc", _pages()[:2], state_dir=state_dir, total=3
    )
    write_logs(legacy)
    _migrate(legacy.root, state_dir)
    legacy.pages = _pages()
    write_logs(legacy)
    _before_apply(
        monkeypatch,
        legacy.root,
        "UPDATE units SET stale = 1 WHERE kind = 'summary' AND idx = 2",
    )

    report = _migrate(legacy.root, state_dir)

    assert _changed(report)
    units = {(u["kind"], u["idx"]): u for u in _view(legacy.root, "units")}
    assert (units[("summary", 2)]["status"], units[("summary", 2)]["stale"]) == (
        "pending",
        1,
    )


def test_a_summary_whose_transcription_a_run_changed_is_left_and_reported(
    tmp_path: Path, state_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    legacy = make_item(tmp_path / "out", "doc", _pages(), state_dir=state_dir)
    write_logs(legacy)
    summary_log = legacy.paths.legacy_logs()[1]
    lines = summary_log.read_text(encoding="utf-8").splitlines(keepends=True)
    summary_log.write_text("".join(lines[:-1]), encoding="utf-8")
    _migrate(legacy.root, state_dir)
    connection = sqlite3.connect(legacy.root / DATABASE_NAME)
    with connection:
        connection.execute(
            "INSERT INTO runs (run_id, tool, tool_version, command, started_at) "
            "VALUES ('a-run', 'autoexcerpter', '2.0.0', 'run', '2026')"
        )
        connection.execute(
            "UPDATE units SET run_id = 'a-run' WHERE kind = 'transcription' AND idx = 2"
        )
    connection.close()
    write_logs(legacy)
    _before_apply(
        monkeypatch,
        legacy.root,
        "UPDATE units SET content = 'Again.', updated_at = '2027' "
        "WHERE kind = 'transcription' AND idx = 2",
    )

    report = _migrate(legacy.root, state_dir)

    assert _changed(report)
    units = {(u["kind"], u["idx"]): u for u in _view(legacy.root, "units")}
    assert units[("summary", 2)]["status"] == "pending"


def test_a_changed_log_transcription_imports_its_summary_again(
    item: Item, state_dir: Path
) -> None:
    _migrate(item.root, state_dir)
    item.pages[0].text = "First page, corrected."
    write_logs(item)

    report = _migrate(item.root, state_dir)

    (document,) = report.documents
    assert document.imported == 2
    units = {(u["kind"], u["idx"]): u for u in _view(item.root, "units")}
    assert units[("transcription", 0)]["content"].endswith("First page, corrected.")
    assert (units[("summary", 0)]["status"], units[("summary", 0)]["stale"]) == (
        "done",
        0,
    )


def test_a_unit_the_store_lacks_is_a_mismatch(
    item: Item, state_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    write = migrate_module._write_document

    def dropping(connection: sqlite3.Connection, changes: Any, *args: Any) -> None:
        changes.transcriptions = changes.transcriptions[1:]
        write(connection, changes, *args)

    monkeypatch.setattr(migrate_module, "_write_document", dropping)

    report = _migrate(item.root, state_dir)

    assert report.exit_code == 1
    assert any("not in the store" in m for m in report.documents[0].mismatches)


def test_a_partial_item_plans_its_missing_pages(
    tmp_path: Path, state_dir: Path
) -> None:
    legacy = make_item(
        tmp_path / "out", "doc", _pages()[:2], state_dir=state_dir, total=4
    )
    write_logs(legacy)

    report = _migrate(legacy.root, state_dir)

    assert report.exit_code == 0
    units = {(u["kind"], u["idx"]): u["status"] for u in _view(legacy.root, "units")}
    assert units[("transcription", 3)] == "pending"
    assert units[("summary", 2)] == "pending"
    assert report.documents[0].after == "partial"


def test_failed_pages_keep_their_status(tmp_path: Path, state_dir: Path) -> None:
    pages = [Page("Ok."), Page(error="timeout"), Page("Ok.", summary_error="refused")]
    legacy = make_item(tmp_path / "out", "doc", pages, state_dir=state_dir)
    write_logs(legacy)
    write_files(legacy)

    report = _migrate(legacy.root, state_dir)

    assert report.exit_code == 0
    (document,) = report.documents
    assert document.failures == 2
    assert document.warnings == []
    summaries = _view(legacy.root, "page_summaries")
    assert [row["outcome"] for row in summaries] == ["ok", "placeholder", "failed"]
    assert [row["status"] for row in _view(legacy.root, "page_text")] == [
        "done",
        "failed",
        "done",
    ]


def test_a_truncated_log_tail_is_dropped_with_a_warning(
    tmp_path: Path, state_dir: Path
) -> None:
    legacy = make_item(tmp_path / "out", "doc", _pages(), state_dir=state_dir)
    write_logs(legacy, tail='{"original_input_order_index": 2, "bull')

    report = _migrate(legacy.root, state_dir)

    assert report.exit_code == 0
    assert any("truncated" in w for w in report.documents[0].warnings)


def test_a_log_without_the_version_2_header_is_an_error(
    item: Item, state_dir: Path
) -> None:
    log = item.paths.legacy_logs()[0]
    log.write_text('{"_format_version": 1}\n', encoding="utf-8")

    report = _migrate(item.root, state_dir)

    assert report.exit_code == 1
    assert report.documents == []
    assert "_format_version" in report.errors[0]


def test_a_dry_run_writes_nothing(item: Item, state_dir: Path) -> None:
    report = _migrate(item.root, state_dir, dry_run=True)

    assert report.documents[0].imported == 6
    assert report.databases[0].action == "create"
    assert not (item.root / DATABASE_NAME).exists()


@pytest.mark.parametrize(
    ("rendered", "existing", "matches"),
    [
        ("# h\n\nA\n", "# old\n# x\n\nA", True),
        ('# h\n\n<page_break pdf="1"/>\nA\n', "# old\n\nA\n", True),
        ("# h\n\nA\n", "# old\n\nA\n\n", False),
        ("# h\n\nA\n", "# old\n\nB\n", False),
    ],
)
def test_bodies_are_compared_without_header_markers_and_final_newline(
    rendered: str, existing: str, matches: bool
) -> None:
    assert body_matches("transcription-md", rendered, existing) is matches


def test_the_summary_metadata_line_is_left_out() -> None:
    old = "# doc\n\n*Summary generated 2026-01-01 10:00 · 1 content pages*\n\nA\n"
    new = "# doc\n\n*Summary generated 2026-10-11 12:00 · 1 content pages*\n\nA\n"

    assert body_matches("summary-md", new, old)
