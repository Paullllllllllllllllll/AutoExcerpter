"""Round 3: rendering and LLM-parameter regression tests.

One focused regression per verified fix:

1. Inline math (``$...$``) lands as a bare ``m:oMath`` interleaved with the
   surrounding runs; only display math (``$$...$$``) gets the block-level
   ``m:oMathPara`` wrapper that Word renders on its own centered line.
2. A rejected currency capture no longer swallows the opening dollar of a
   genuine formula that follows it (odd total dollar count).
3. Non-string debris in ``page_types`` (a corrupt or reused log) no longer
   raises ``TypeError`` inside the writers' shared preparation path.
4. The text writer disables platform newline translation, so a transcription
   containing a literal CRLF is written verbatim instead of as ``\\r\\r\\n``.
5. Adaptive-thinking Claude models (4.6+) receive ``output_config.effort``
   instead of ``thinking.budget_tokens``, which they reject with HTTP 400.
6. File-specific summary context for an image folder resolves against the
   FULL directory name (mirroring ``ItemSpec.output_stem``), not the
   ``Path.stem``-collapsed name.
7. A markdown-fenced plain-text sentinel still triggers the configured
   sentinel retries (the parser already recognized fenced sentinels; the
   retry gate compared the raw text and skipped them).
"""

from __future__ import annotations

import threading
from collections import deque
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock, patch

from docx import Document

from llm.base import LLMClientBase
from llm.client import get_model_capabilities
from llm.transcription import TranscriptionManager
from pipeline.context import resolve_summary_context
from rendering.citations import CitationManager
from rendering.docx import add_formatted_text_to_paragraph, parse_latex_in_text
from rendering.summary import prepare_summary_data
from rendering.text import write_transcription_to_text


# ---------------------------------------------------------------------------
# Fix 1: inline math must not be wrapped in the display-level m:oMathPara
# ---------------------------------------------------------------------------
class TestInlineMathStaysInline:
    """Inline formulas flow with the text; display formulas keep the wrapper."""

    def test_inline_formula_is_bare_omath_between_runs(self) -> None:
        paragraph = Document().add_paragraph()
        add_formatted_text_to_paragraph(paragraph, "The value $x^2$ is small.")

        xml = paragraph._p.xml
        assert "oMath" in xml
        assert "oMathPara" not in xml
        # The surrounding prose is preserved as ordinary runs around the math.
        texts = [run.text for run in paragraph.runs]
        assert "The value " in texts
        assert " is small." in texts

    def test_display_formula_keeps_omathpara_wrapper(self) -> None:
        paragraph = Document().add_paragraph()
        add_formatted_text_to_paragraph(paragraph, "$$E = mc^2$$")

        assert "oMathPara" in paragraph._p.xml

    def test_mixed_line_wraps_only_the_display_formula(self) -> None:
        paragraph = Document().add_paragraph()
        add_formatted_text_to_paragraph(
            paragraph, "Given $x + y$ we obtain $$E = mc^2$$ at last."
        )

        xml = paragraph._p.xml
        assert xml.count("oMathPara") == 2  # one open + one close tag


# ---------------------------------------------------------------------------
# Fix 2: currency guard must not consume a following formula's opening dollar
# ---------------------------------------------------------------------------
class TestCurrencyGuardResume:
    """With an odd dollar count the trailing formula must still parse."""

    def test_math_after_currency_with_odd_dollar_count(self) -> None:
        segments = parse_latex_in_text("cost of $5 and weight of $x^2$")

        assert ("x^2", "latex_inline") in segments
        joined = "".join(content for content, kind in segments if kind == "text")
        assert "$5" in joined

    def test_currency_only_line_stays_literal(self) -> None:
        segments = parse_latex_in_text("Wages rose from $3 to $5 that year.")

        assert all(kind == "text" for _, kind in segments)

    def test_consecutive_currency_before_math(self) -> None:
        segments = parse_latex_in_text("price $5, variable $x$ and $y$")

        assert ("x", "latex_inline") in segments
        assert ("y", "latex_inline") in segments
        assert not any(
            kind == "latex_inline" and content.strip() in {"and", ", variable"}
            for content, kind in segments
        )


# ---------------------------------------------------------------------------
# Fix 3: non-string page_types entries must not crash the render
# ---------------------------------------------------------------------------
class TestPageTypesDebrisTolerated:
    """Unhashable page_types entries from a corrupt log are dropped."""

    @staticmethod
    def _result(page_types: Any) -> dict[str, Any]:
        return {
            "page": 1,
            "page_information": {
                "page_number_integer": 1,
                "page_number_type": "arabic",
                "page_types": page_types,
                "is_two_page_spread": False,
            },
            "bullet_points": ["A finding."],
            "references": [],
        }

    def test_dict_entry_does_not_raise(self) -> None:
        data = prepare_summary_data(
            [self._result([{"bad": True}])], CitationManager()
        )
        # The malformed entry degrades to the "content" default; the page
        # still renders with its bullet intact.
        assert len(data.page_render_items) == 1

    def test_mixed_entries_keep_string_types(self) -> None:
        data = prepare_summary_data(
            [self._result(["bibliography", {"bad": True}])], CitationManager()
        )
        assert data.page_type_pages.get("bibliography") == [(1, "arabic")]


# ---------------------------------------------------------------------------
# Fix 4: text writer preserves CRLF verbatim (no \r\r\n mojibake)
# ---------------------------------------------------------------------------
class TestTextWriterNewlines:
    """newline="\\n" disables platform newline translation on write."""

    def test_crlf_in_transcription_survives_verbatim(self, tmp_path: Path) -> None:
        out = tmp_path / "book.txt"
        ok = write_transcription_to_text(
            transcription_results=[{"transcription": "line1\r\nline2"}],
            output_path=out,
            document_name="book",
            item_type="PDF",
            total_elapsed_time=1.0,
            source_path=tmp_path / "book.pdf",
        )

        assert ok is True
        data = out.read_bytes()
        assert b"\r\r\n" not in data
        assert b"line1\r\nline2" in data


# ---------------------------------------------------------------------------
# Fix 5: adaptive-thinking Claude models must not receive budget_tokens
# ---------------------------------------------------------------------------
def _make_client(**overrides: Any) -> LLMClientBase:
    """Create a bare LLMClientBase bypassing __init__ for isolated unit tests."""
    client = LLMClientBase.__new__(LLMClientBase)
    defaults: dict[str, Any] = {
        "model_name": "gpt-5-mini",
        "provider": "openai",
        "timeout": 300,
        "rate_limiter": None,
        "max_retries": 5,
        "chat_model": MagicMock(),
        "_stats_lock": threading.Lock(),
        "successful_requests": 0,
        "failed_requests": 0,
        "processing_times": deque(maxlen=50),
        "model_config": {},
        "service_tier": "auto",
        "schema_retry_config": {},
        "_output_schema": None,
    }
    defaults.update(overrides)
    for attr, val in defaults.items():
        setattr(client, attr, val)
    return client


_ADAPTIVE_CAPS = {
    "max_tokens": True,
    "reasoning": False,
    "extended_thinking": False,
    "adaptive_thinking": True,
    "thinking": False,
    "text_verbosity": False,
}


class TestAnthropicAdaptiveThinking:
    """Claude 4.6+ takes output_config.effort; budget_tokens would be HTTP 400."""

    def test_capability_split_adaptive_vs_budget_generation(self) -> None:
        adaptive = get_model_capabilities("claude-fable-5")
        assert adaptive["adaptive_thinking"] is True
        assert adaptive["extended_thinking"] is False

        legacy = get_model_capabilities("claude-opus-4-5")
        assert legacy["adaptive_thinking"] is False
        assert legacy["extended_thinking"] is True

    @patch("llm.base.get_model_capabilities", return_value=_ADAPTIVE_CAPS)
    def test_adaptive_model_gets_effort_not_budget(self, _: Any) -> None:
        client = _make_client(
            provider="anthropic",
            model_name="claude-fable-5",
            model_config={"reasoning": {"effort": "high"}},
            service_tier="",
        )
        kwargs = client._build_invoke_kwargs()
        assert kwargs["output_config"] == {"effort": "high"}
        assert "thinking" not in kwargs

    @patch("llm.base.get_model_capabilities", return_value=_ADAPTIVE_CAPS)
    def test_adaptive_model_maps_minimal_to_low(self, _: Any) -> None:
        client = _make_client(
            provider="anthropic",
            model_name="claude-sonnet-5",
            model_config={"reasoning": {"effort": "minimal"}},
            service_tier="",
        )
        kwargs = client._build_invoke_kwargs()
        assert kwargs["output_config"] == {"effort": "low"}
        assert "thinking" not in kwargs

    @patch("llm.base.get_model_capabilities", return_value=_ADAPTIVE_CAPS)
    def test_adaptive_model_effort_none_sends_nothing(self, _: Any) -> None:
        """"none" mirrors the budget path: no reasoning parameter at all."""
        client = _make_client(
            provider="anthropic",
            model_name="claude-fable-5",
            model_config={"reasoning": {"effort": "none"}},
            service_tier="",
        )
        kwargs = client._build_invoke_kwargs()
        assert "output_config" not in kwargs
        assert "thinking" not in kwargs


# ---------------------------------------------------------------------------
# Fix 6: folder-specific context resolves against the full directory name
# ---------------------------------------------------------------------------
class TestFolderContextFullName:
    """A dotted image-folder name must find its own context file."""

    def test_dotted_folder_finds_its_context_file(self, tmp_path: Path) -> None:
        folder = tmp_path / "photos.2023"
        folder.mkdir()
        ctx = tmp_path / "photos.2023_summary_context.txt"
        ctx.write_text("Context for the 2023 photo set.", encoding="utf-8")
        # A sibling document whose stem collides with the collapsed name: the
        # old Path.stem lookup would have picked THIS file for the folder.
        wrong = tmp_path / "photos_summary_context.txt"
        wrong.write_text("Context for photos.pdf, not the folder.", encoding="utf-8")

        content, resolved = resolve_summary_context(
            folder, global_context_dir=tmp_path / "no_such_ctx_dir"
        )

        assert resolved == ctx
        assert content == "Context for the 2023 photo set."

    def test_pdf_input_still_uses_stem(self, tmp_path: Path) -> None:
        pdf = tmp_path / "book.pdf"
        pdf.write_bytes(b"%PDF-1.4")
        ctx = tmp_path / "book_summary_context.txt"
        ctx.write_text("Context for the book.", encoding="utf-8")

        content, resolved = resolve_summary_context(
            pdf, global_context_dir=tmp_path / "no_such_ctx_dir"
        )

        assert resolved == ctx
        assert content == "Context for the book."


# ---------------------------------------------------------------------------
# Fix 7: fenced plain-text sentinels still trigger sentinel retries
# ---------------------------------------------------------------------------
def _make_plain_text_tm(**overrides: Any) -> TranscriptionManager:
    """Bare plain-text-mode TranscriptionManager bypassing __init__."""
    tm = TranscriptionManager.__new__(TranscriptionManager)
    caps = MagicMock()
    caps.use_plain_text_prompt = True
    defaults: dict[str, Any] = {
        "model_name": "gpt-5-mini",
        "provider": "openai",
        "custom_capabilities": caps,
        "_stats_lock": threading.Lock(),
        "successful_requests": 0,
        "failed_requests": 0,
        "processing_times": deque(maxlen=50),
        "rate_limiter": None,
        "model_config": {},
        "system_prompt": "sys",
        "_output_schema": None,
        "schema_retry_config": {
            "no_transcribable_text": {
                "enabled": True,
                "max_attempts": 2,
                "backoff_base": 0.0,
                "backoff_multiplier": 1.0,
            }
        },
    }
    defaults.update(overrides)
    for attr, val in defaults.items():
        setattr(tm, attr, val)
    return tm


def _wire_tm(tm: TranscriptionManager, extract: Any) -> None:
    """Stub the API boundary; *extract* is the return_value or side_effect."""
    tm._build_model_inputs = MagicMock(return_value=([], {}))  # type: ignore[method-assign]
    tm._get_structured_chat_model = MagicMock(return_value=MagicMock())  # type: ignore[method-assign]
    tm._invoke_with_retry = MagicMock(return_value=MagicMock())  # type: ignore[method-assign]
    tm._report_token_usage = MagicMock()  # type: ignore[method-assign]
    mock = MagicMock()
    if isinstance(extract, list):
        mock.side_effect = extract
    else:
        mock.return_value = extract
    tm._extract_output_text = mock  # type: ignore[method-assign]


def _payload(name: str = "page_001.png") -> Any:
    payload = MagicMock()
    payload.image_name = name
    payload.sequence_number = 1
    payload.base64 = "AAAA"
    return payload


class TestFencedSentinelRetries:
    """A code-fenced sentinel must consume a sentinel retry like a bare one."""

    def test_fenced_sentinel_retries_then_succeeds(self) -> None:
        tm = _make_plain_text_tm()
        _wire_tm(
            tm,
            ["```\n[no transcribable text]\n```", "Real transcribed text."],
        )

        with patch("llm.transcription.time.sleep"):
            result = tm.transcribe_payload(_payload(), max_schema_retries=3)

        assert result["transcription"] == "Real transcribed text."
        assert result["schema_retries"]["no_transcribable_text"] == 1

    def test_bare_sentinel_retry_behavior_unchanged(self) -> None:
        tm = _make_plain_text_tm()
        _wire_tm(tm, ["[no transcribable text]", "Recovered text."])

        with patch("llm.transcription.time.sleep"):
            result = tm.transcribe_payload(_payload(), max_schema_retries=3)

        assert result["transcription"] == "Recovered text."
        assert result["schema_retries"]["no_transcribable_text"] == 1
