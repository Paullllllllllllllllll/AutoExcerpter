"""Round 5: DOCX/Markdown writer regression tests.

One focused regression per verified fix:

1. ``add_math_to_paragraph`` catches lxml's ``XMLSyntaxError`` (a
   ``SyntaxError`` subclass), so the OMML sanitize-and-retry path is reachable
   instead of every malformed equation degrading to italic text.
2. Page headings are XML-sanitized in both writers, so an illegal control
   character in a page number no longer fails the whole document.
3. Both writers clean up their ``.tmp`` sibling on non-``OSError`` failures.
4. Markdown link URLs escape angle brackets, and the file ends with a newline.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest
from docx import Document

import rendering.docx as docx_module
import rendering.markdown as markdown_module
from rendering.citations import CitationManager
from rendering.docx import add_math_to_paragraph, create_docx_summary
from rendering.markdown import create_markdown_summary
from rendering.summary import prepare_summary_data

_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# mathml2omml occasionally emits a <m:groupChrPr> that is closed by the parent
# </m:groupChr>; lxml rejects it with XMLSyntaxError.
_BROKEN_OMML = (
    f'<m:oMath xmlns:m="{_MATH_NS}"><m:groupChr><m:groupChrPr>'
    '<m:chr m:val="x"/><m:e><m:r><m:t>a</m:t></m:r></m:e>'
    "</m:groupChr></m:oMath>"
)
_REPAIRED_OMML = f'<m:oMath xmlns:m="{_MATH_NS}"><m:r><m:t>a</m:t></m:r></m:oMath>'

# Malformed markup with no groupChrPr construct: the sanitizer has no repair
# for it, so the italic-text fallback must fire.
_UNREPAIRABLE_OMML = f'<m:oMath xmlns:m="{_MATH_NS}"><m:r><m:t>a</m:r></m:t></m:oMath>'


def _content_page(
    page_number: Any, bullets: list[str], references: Any = None
) -> dict[str, Any]:
    """Build a minimal flat summary result for a content page."""
    return {
        "page": page_number,
        "page_information": {
            "page_number_integer": page_number,
            "is_two_page_spread": False,
            "page_number_integer_end": None,
            "page_number_type": "arabic",
            "page_types": ["content"],
        },
        "bullet_points": bullets,
        "references": references,
    }


# ---------------------------------------------------------------------------
# Fix 1: lxml XMLSyntaxError reaches the OMML sanitize branch
# ---------------------------------------------------------------------------
class TestOmmlSanitizeBranchReachable:
    """A parse failure from python-docx must route into sanitize_omml_xml."""

    def test_broken_omml_raises_syntax_error_not_value_error(self) -> None:
        """Pins the premise: lxml's error is a SyntaxError, not a ValueError."""
        from docx.oxml import parse_xml

        with pytest.raises(SyntaxError):
            parse_xml(f'<m:oMathPara xmlns:m="{_MATH_NS}">{_BROKEN_OMML}</m:oMathPara>')

    def test_sanitized_markup_is_used_instead_of_italic_fallback(self) -> None:
        """The recovery path runs and native OMML lands in the paragraph."""
        paragraph = Document().add_paragraph()
        seen: list[str] = []

        def fake_sanitize(markup: str) -> str:
            seen.append(markup)
            return _REPAIRED_OMML

        with (
            patch("rendering.docx.mathml2omml.convert", return_value=_BROKEN_OMML),
            patch.object(docx_module, "sanitize_omml_xml", side_effect=fake_sanitize),
        ):
            add_math_to_paragraph(paragraph, "x^2")

        assert len(seen) == 1
        xml = paragraph._p.xml
        assert "oMathPara" in xml
        assert "Cambria Math" not in xml
        assert paragraph.runs == []

    def test_real_sanitizer_repairs_groupchr_markup_end_to_end(self) -> None:
        """The real sanitizer fixes the missing </m:groupChrPr> close tag."""
        paragraph = Document().add_paragraph()

        with patch("rendering.docx.mathml2omml.convert", return_value=_BROKEN_OMML):
            add_math_to_paragraph(paragraph, "x^2")

        xml = paragraph._p.xml
        assert "oMathPara" in xml
        assert "Cambria Math" not in xml
        assert paragraph.runs == []

    def test_unrepairable_markup_still_degrades_to_italic_text(self) -> None:
        """Markup the sanitizer cannot fix falls back to italic text."""
        paragraph = Document().add_paragraph()

        with patch(
            "rendering.docx.mathml2omml.convert", return_value=_UNREPAIRABLE_OMML
        ):
            add_math_to_paragraph(paragraph, "x^2")

        assert paragraph.runs
        assert paragraph.runs[0].italic is True
        assert paragraph.runs[0].font.name == "Cambria Math"


# ---------------------------------------------------------------------------
# Fix 2: page headings are XML-sanitized in both writers
# ---------------------------------------------------------------------------
class TestHeadingSanitized:
    """A control character in a page number must not fail the document."""

    _RESULT = _content_page("7\x0b", ["A bullet"])

    def test_heading_carries_the_illegal_character(self) -> None:
        """Pins the premise: the heading really does contain the control char."""
        data = prepare_summary_data([self._RESULT], CitationManager())
        assert data.page_render_items[0].heading_text == "Page 7\x0b"

    def test_docx_renders_instead_of_raising(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.docx"
        with patch.object(docx_module, "enrich_if_enabled"):
            create_docx_summary([self._RESULT], output_path, "Doc")

        assert output_path.exists()
        texts = [p.text for p in Document(str(output_path)).paragraphs]
        assert "Page 7" in texts
        assert not any("\x0b" in t for t in texts)

    def test_markdown_renders_sanitized_heading(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.md"
        with patch.object(markdown_module, "enrich_if_enabled"):
            create_markdown_summary([self._RESULT], output_path, "Doc")

        content = output_path.read_text(encoding="utf-8")
        assert "### Page 7" in content
        assert "\x0b" not in content


# ---------------------------------------------------------------------------
# Fix 3: non-OSError failures still clean up the .tmp sibling
# ---------------------------------------------------------------------------
class TestNonOsErrorTempCleanup:
    """A ValueError mid-write must not orphan the temp file."""

    _RESULT = _content_page(1, ["A bullet"])

    @staticmethod
    def _boom(_src: Any, _dst: Any) -> None:
        raise ValueError("not an OSError")

    def test_docx_writer_removes_tmp(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        output_path = tmp_path / "summary.docx"
        monkeypatch.setattr("rendering.docx.os.replace", self._boom)

        with (
            patch.object(docx_module, "enrich_if_enabled"),
            pytest.raises(ValueError, match="not an OSError"),
        ):
            create_docx_summary([self._RESULT], output_path, "Doc")

        assert not output_path.with_name(output_path.name + ".tmp").exists()
        assert [p.name for p in tmp_path.iterdir() if p.suffix == ".tmp"] == []

    def test_markdown_writer_removes_tmp(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        output_path = tmp_path / "summary.md"
        monkeypatch.setattr("rendering.markdown.os.replace", self._boom)

        with (
            patch.object(markdown_module, "enrich_if_enabled"),
            pytest.raises(ValueError, match="not an OSError"),
        ):
            create_markdown_summary([self._RESULT], output_path, "Doc")

        assert not output_path.with_name(output_path.name + ".tmp").exists()
        assert [p.name for p in tmp_path.iterdir() if p.suffix == ".tmp"] == []


# ---------------------------------------------------------------------------
# Fix 4: markdown link escaping and trailing newline
# ---------------------------------------------------------------------------
class TestMarkdownLinkAndNewline:
    """Angle brackets in a URL must be escaped; the file must end in a newline."""

    _REFERENCED = _content_page(
        1,
        ["A bullet"],
        references=[
            {
                "citation": "Smith, J. (2020). *A Linked Work*. Press.",
                "is_partial": False,
            }
        ],
    )

    def _render_with_url(self, tmp_path: Path, url: str) -> str:
        manager = CitationManager()
        data = prepare_summary_data([self._REFERENCED], manager)
        for citation in manager.citations.values():
            citation.url = url

        output_path = tmp_path / "summary.md"
        create_markdown_summary(
            [self._REFERENCED],
            output_path,
            "Doc",
            citation_manager=manager,
            data=data,
        )
        return output_path.read_text(encoding="utf-8")

    def test_angle_brackets_escaped_in_link_url(self, tmp_path: Path) -> None:
        content = self._render_with_url(
            tmp_path, "https://doi.org/10.1000/j.<seg>.2020"
        )
        assert "%3Cseg%3E" in content
        assert "<seg>" not in content

    def test_parentheses_still_escaped(self, tmp_path: Path) -> None:
        content = self._render_with_url(
            tmp_path, "https://doi.org/10.1016/S0140-6736(00)02800-3"
        )
        assert "%28" in content and "%29" in content

    def test_file_ends_with_newline(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.md"
        with patch.object(markdown_module, "enrich_if_enabled"):
            create_markdown_summary(
                [_content_page(1, ["A bullet"])], output_path, "Doc"
            )

        content = output_path.read_text(encoding="utf-8")
        assert content.endswith("\n")
        assert not content.endswith("\n\n\n")

    def test_empty_document_ends_with_newline(self, tmp_path: Path) -> None:
        output_path = tmp_path / "empty.md"
        with patch.object(markdown_module, "enrich_if_enabled"):
            create_markdown_summary([], output_path, "Empty Doc")

        assert output_path.read_text(encoding="utf-8").endswith("\n")
