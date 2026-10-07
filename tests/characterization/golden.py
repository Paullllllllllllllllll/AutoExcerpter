"""Normalized golden snapshots of a characterization run.

``snapshot`` turns a ``RunResult`` into text files: every file the run
created (the transcription, the summary ``.md``, text dumps of the summary
``.docx`` and of the SQLite database, the JSONL working logs), the JSON
summary line with the exit code, and ``requests.json`` with the
provider-class constructions, the per-call request digests and the OpenAlex
requests. Paths, dates, timestamps, durations, hashes and versions become
placeholders. ``check_golden`` compares the snapshot with ``golden/<case>/``
and fails with a unified diff, or rewrites the case when
``AE_UPDATE_GOLDEN=1``.
"""

from __future__ import annotations

import contextlib
import difflib
import json
import os
import re
import shutil
from collections.abc import Callable
from pathlib import Path
from typing import Any

import docx
from docx.text.paragraph import Paragraph

from autoexcerpter.common.sqlite import read_rows, table_names
from autoexcerpter.rendering.sqlite import TABLES
from tests.characterization.adapter import RunResult
from tests.characterization.fakes import FakeOpenAlex

GOLDEN_DIR = Path(__file__).resolve().parent / "golden"
UPDATE_ENV = "AE_UPDATE_GOLDEN"

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

_DATETIME_RE = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?")
_DURATION_RE = re.compile(r"(Total processing time for this item: )[^\n]*")
_TMP_PATH_RE = re.compile(r"<tmp>[^\s\"'<>|*?\]]*")

_KEY_PLACEHOLDERS: dict[str, str] = {
    "processing_time": "<duration>",
    "elapsed_s": "<duration>",
    "avg_api_s": "<duration>",
    "processing_start_time": "<datetime>",
    "source_sha256": "<sha256>",
    "image_names_sha256": "<sha256>",
    "pymupdf_version": "<version>",
    "pillow_version": "<version>",
    "byte_size": "<bytes>",
    "size": "<bytes>",
    "total_image_bytes": "<bytes>",
}
# Columns of the SQLite dump that vary between identical runs.
_COLUMN_PLACEHOLDERS: dict[str, str] = {
    "hash": "<sha256>",
    "tool_version": "<version>",
}
_TABLE_KEYS = {spec.name: list(spec.key) for spec in TABLES}


# ============================================================================
# Normalization
# ============================================================================
class Normalizer:
    """Replace run-specific values with stable placeholders."""

    def __init__(self, root: Path) -> None:
        forms: set[str] = set()
        for candidate in (root, root.resolve()):
            text = str(candidate)
            forms.update({text, text.replace("\\", "/"), text.replace("/", "\\")})
        self._roots = sorted(forms, key=len, reverse=True)

    def text(self, value: str) -> str:
        """Normalize paths, dates and durations in free text."""
        for form in self._roots:
            value = value.replace(form, "<tmp>")
        value = _TMP_PATH_RE.sub(lambda m: m.group(0).replace("\\", "/"), value)
        value = _DATETIME_RE.sub("<datetime>", value)
        return _DURATION_RE.sub(r"\1<duration>", value)

    def data(self, value: Any, key: str | None = None) -> Any:
        """Normalize a parsed JSON value recursively."""
        if key in _KEY_PLACEHOLDERS and value is not None:
            return _KEY_PLACEHOLDERS[key]
        if isinstance(value, dict):
            return {k: self.data(v, k) for k, v in value.items()}
        if isinstance(value, list):
            return [self.data(item) for item in value]
        if isinstance(value, str):
            return self.text(value)
        return value


def _dumps(value: Any) -> str:
    return json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n"


def _jsonl(path: Path, norm: Normalizer) -> str:
    """Normalize a JSONL working log: header first, entries in page order."""
    header: list[Any] = []
    entries: list[Any] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        record = json.loads(line)
        if isinstance(record, dict) and "_format_version" in record:
            header.append(record)
        else:
            entries.append(record)

    def order(entry: Any) -> tuple[int, str]:
        if not isinstance(entry, dict):
            return 10_000, ""
        index = entry.get("original_input_order_index")
        return (index if isinstance(index, int) else 10_000), str(entry.get("image"))

    entries.sort(key=order)
    return "".join(_dumps(norm.data(record)) for record in header + entries)


def _run_text(run: Any) -> str:
    parts: list[str] = []
    for child in run.iterchildren():
        if child.tag in (_W + "t", _W + "instrText"):
            parts.append(child.text or "")
        elif child.tag == _W + "tab":
            parts.append("\t")
        elif child.tag == _W + "br":
            kind = child.get(_W + "type")
            parts.append("[page-break]" if kind == "page" else "\n")
        elif child.tag == _W + "fldChar":
            parts.append(f"[field-{child.get(_W + 'fldCharType')}]")
    return "".join(parts)


def _flag(run: Any, name: str) -> bool:
    rpr = run.find(_W + "rPr")
    element = rpr.find(_W + name) if rpr is not None else None
    if element is None:
        return False
    return element.get(_W + "val") not in ("0", "false", "off")


def _styled(run: Any) -> str:
    text = _run_text(run)
    if not text:
        return ""
    if _flag(run, "i"):
        text = f"<i>{text}</i>"
    if _flag(run, "b"):
        text = f"<b>{text}</b>"
    return text


def _math_text(element: Any) -> str:
    return "".join(node.text or "" for node in element.iter(_M + "t"))


def _paragraph_content(p: Any, rels: Any) -> str:
    parts: list[str] = []
    for child in p.iterchildren():
        if child.tag == _W + "r":
            parts.append(_styled(child))
        elif child.tag == _W + "hyperlink":
            rel_id = child.get(_R + "id")
            target = rels[rel_id].target_ref if rel_id in rels else "?"
            text = "".join(_styled(run) for run in child.iter(_W + "r"))
            parts.append(f'<a href="{target}">{text}</a>')
        elif child.tag == _M + "oMath":
            parts.append(f"<math>{_math_text(child)}</math>")
        elif child.tag == _M + "oMathPara":
            parts.append(f"<mathpara>{_math_text(child)}</mathpara>")
    return "".join(parts)


def dump_docx(path: Path) -> str:
    """Return a text dump of a DOCX body: styles, runs, links, tables, math."""
    document = docx.Document(str(path))
    body = document.element.body
    rels = document.part.rels
    lines: list[str] = []
    for child in body.iterchildren():
        if child.tag == _W + "p":
            style = Paragraph(child, document._body).style
            name = style.name if style is not None else "?"
            lines.append(f"[{name}] {_paragraph_content(child, rels)}".rstrip())
        elif child.tag == _W + "tbl":
            rows = child.findall(_W + "tr")
            cols = max((len(row.findall(_W + "tc")) for row in rows), default=0)
            lines.append(f"[Table {len(rows)}x{cols}]")
            for row in rows:
                cells = [
                    " / ".join(_paragraph_content(p, rels) for p in cell.iter(_W + "p"))
                    for cell in row.findall(_W + "tc")
                ]
                lines.append("| " + " | ".join(cells) + " |")
    lines.append(f"[oMath elements: {len(body.findall('.//' + _M + 'oMath'))}]")
    return "\n".join(lines) + "\n"


# ============================================================================
# Snapshot and comparison
# ============================================================================
def _column_value(column: str, value: Any, norm: Normalizer) -> Any:
    if column in _COLUMN_PLACEHOLDERS and value is not None:
        return _COLUMN_PLACEHOLDERS[column]
    if isinstance(value, str) and value[:1] in "[{":
        with contextlib.suppress(json.JSONDecodeError):
            value = json.loads(value)
    return norm.data(value, column)


def dump_sqlite(path: Path, norm: Normalizer) -> str:
    """Return every table of a SQLite file, rows ordered by key, normalized.

    JSON text columns are parsed so that their paths normalize too.
    """
    parts: list[str] = []
    for table in table_names(path):
        parts.append(f"## {table}\n")
        for row in read_rows(path, table, order_by=_TABLE_KEYS.get(table)):
            parts.append(
                _dumps({key: _column_value(key, val, norm) for key, val in row.items()})
            )
    return "".join(parts)


def _file_entry(path: Path, norm: Normalizer) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return ".dump.txt", norm.text(dump_docx(path))
    if suffix == ".sqlite":
        return ".dump.txt", dump_sqlite(path, norm)
    if suffix in (".json", ".jsonl"):
        return "", _jsonl(path, norm)
    if suffix in (".txt", ".md"):
        return "", norm.text(path.read_text(encoding="utf-8"))
    return ".bytes.txt", f"<{path.stat().st_size} bytes>\n"


def snapshot(result: RunResult, openalex: FakeOpenAlex | None = None) -> dict[str, str]:
    """Return the normalized golden files of one run, keyed by relative name."""
    norm = Normalizer(result.root)
    files: dict[str, str] = {}
    for path in result.created:
        relative = norm.text(str(path)).removeprefix("<tmp>/")
        extra, content = _file_entry(path, norm)
        files[f"files/{relative}{extra}"] = content
    files["summary.json"] = _dumps(
        {"exit_code": result.exit_code, "json_line": norm.data(result.summary)}
    )
    requests: dict[str, Any] = {
        "models": sorted(
            result.llm.models,
            key=lambda m: json.dumps(m, sort_keys=True, default=str),
        ),
        "calls": result.calls(),
    }
    if openalex is not None:
        requests["openalex"] = sorted(
            openalex.requests,
            key=lambda r: json.dumps(r, sort_keys=True, ensure_ascii=False),
        )
    files["requests.json"] = _dumps(norm.data(requests))
    return files


def _read_case(case_dir: Path) -> dict[str, str]:
    if not case_dir.is_dir():
        return {}
    return {
        path.relative_to(case_dir).as_posix(): path.read_text(encoding="utf-8")
        for path in sorted(case_dir.rglob("*"))
        if path.is_file()
    }


def _write_case(case_dir: Path, files: dict[str, str]) -> None:
    if case_dir.exists():
        shutil.rmtree(case_dir)
    for name, content in files.items():
        target = case_dir / name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8", newline="\n")


def compare(expected: dict[str, str], actual: dict[str, str]) -> list[str]:
    """Return human-readable differences between two snapshots."""
    problems: list[str] = []
    for name in sorted(set(expected) | set(actual)):
        if name not in actual:
            problems.append(f"missing from run: {name}")
        elif name not in expected:
            problems.append(f"not in golden: {name}")
        elif expected[name] != actual[name]:
            diff = difflib.unified_diff(
                expected[name].splitlines(keepends=True),
                actual[name].splitlines(keepends=True),
                fromfile=f"golden/{name}",
                tofile=f"run/{name}",
            )
            problems.append("".join(diff))
    return problems


def check_golden(
    case: str,
    result: RunResult,
    *,
    openalex: FakeOpenAlex | None = None,
    on_write: Callable[[Path], object] | None = None,
) -> None:
    """Compare a run with ``golden/<case>/``; rewrite it when updating.

    *on_write* receives the case directory after a rewrite (the conftest uses
    it to clear the write guard for the golden tree).
    """
    check_files(case, snapshot(result, openalex), on_write=on_write)


def check_files(
    case: str,
    actual: dict[str, str],
    *,
    on_write: Callable[[Path], object] | None = None,
) -> None:
    """Compare text files keyed by relative name with ``golden/<case>/``."""
    case_dir = GOLDEN_DIR / case
    if os.environ.get(UPDATE_ENV) == "1":
        _write_case(case_dir, actual)
        if on_write is not None:
            on_write(case_dir)
        return
    expected = _read_case(case_dir)
    if not expected:
        raise AssertionError(
            f"No golden files for case {case!r}; run with {UPDATE_ENV}=1 to create."
        )
    problems = compare(expected, actual)
    if problems:
        raise AssertionError(
            f"Golden mismatch for case {case!r} (set {UPDATE_ENV}=1 to accept):\n"
            + "\n".join(problems)
        )
