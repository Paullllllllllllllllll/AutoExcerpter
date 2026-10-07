"""The DOCX summary writer in rendering/docx.py.

Covers whole documents written to disk and read back (styles, page geometry,
footer, headings, bullets, references), malformed input and the atomic save,
hyperlinks, formatted text with LaTeX and the Markdown emphasis parser used for
bullets and citations.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from autoexcerpter.rendering.docx import (
    add_formatted_text_to_paragraph,
    add_hyperlink,
    create_docx_summary,
    parse_markdown_emphasis,
    strip_markdown_emphasis,
)
from autoexcerpter.rendering.summary import build_render_context

_HYPERLINK = qn("w:hyperlink")
_TEXT = qn("w:t")


def _content_page(
    number: int,
    bullets: list[str],
    page_types: list[str] | None = None,
    references: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "page_information": {
            "page_number_integer": number,
            "page_number_type": "arabic",
            "page_types": page_types or ["content"],
        },
        "bullet_points": bullets,
        "references": references or [],
    }


# ============================================================================
# Whole documents
# ============================================================================
class TestCreateDocxSummarySmoke:
    """Real DOCX rendering: write to disk and reopen."""

    def test_docx_styles_geometry_and_footer(self, tmp_path: Path) -> None:
        """A written document carries the Normal style, A4 geometry, 2 cm
        margins, a page-number footer and a prefix-free title."""
        output_path = tmp_path / "summary.docx"
        summary_results = [
            _content_page(12, ["First insight.", "Second insight with $x + y$."])
        ]

        create_docx_summary(summary_results, output_path, "Sample Document")

        doc = Document(str(output_path))

        normal = doc.styles["Normal"]
        assert normal.font.name == "Times New Roman"
        assert normal.font.size == Pt(11)

        section = doc.sections[0]
        assert section.page_width is not None
        assert section.page_height is not None
        assert section.top_margin is not None
        assert section.bottom_margin is not None
        assert section.left_margin is not None
        assert section.right_margin is not None
        assert round(section.page_width.cm, 1) == 21.0
        assert round(section.page_height.cm, 1) == 29.7
        assert round(section.top_margin.cm, 1) == 2.0
        assert round(section.bottom_margin.cm, 1) == 2.0
        assert round(section.left_margin.cm, 1) == 2.0
        assert round(section.right_margin.cm, 1) == 2.0

        assert "PAGE" in section.footer.paragraphs[0]._p.xml

        paragraph_texts = [p.text for p in doc.paragraphs]
        assert "Sample Document" in paragraph_texts
        assert all(not t.startswith("Summary of") for t in paragraph_texts)

        # Heading styles pin an explicit face and carry no theme-font
        # attributes, so Word and LibreOffice render them in Times New Roman
        # rather than the theme's major font.
        for style_name in ("Title", "Heading 1", "Heading 2"):
            rpr = doc.styles[style_name].element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            assert rfonts is not None, f"{style_name} has no rFonts"
            assert rfonts.get(qn("w:ascii")) == "Times New Roman"
            assert rfonts.get(qn("w:hAnsi")) == "Times New Roman"
            for theme_attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                assert rfonts.get(qn(f"w:{theme_attr}")) is None, (
                    f"{style_name} retains w:{theme_attr}"
                )

    def test_heading_and_bullet_style_ids(self, tmp_path: Path) -> None:
        """Title, Document Structure, page headings and bullets carry the
        Title, Heading 1, Heading 2 and List Bullet styles."""
        output_path = tmp_path / "summary.docx"
        summary_results = [
            _content_page(1, ["Point A", "Point B"]),
            _content_page(50, [], page_types=["bibliography"]),
        ]

        create_docx_summary(summary_results, output_path, "Test")

        doc = Document(str(output_path))
        style_by_text = {
            p.text: p.style.name for p in doc.paragraphs if p.style is not None
        }

        assert style_by_text["Test"] == "Title"
        assert style_by_text["Document Structure"] == "Heading 1"
        assert style_by_text["Page 1"] == "Heading 2"

        # A bibliography page adds a page-range line in the Normal style.
        structure_line = next(
            p for p in doc.paragraphs if p.text.startswith("Bibliography:")
        )
        assert structure_line.text == "Bibliography: p. 50"
        assert structure_line.style is not None
        assert structure_line.style.name == "Normal"

        bullets = [
            p
            for p in doc.paragraphs
            if p.style is not None and p.style.name == "List Bullet"
        ]
        assert [p.text for p in bullets] == ["Point A", "Point B"]

    def test_citation_markdown_emphasis_rendered_as_runs(self, tmp_path: Path) -> None:
        """Markdown emphasis in citation text becomes italic runs, not
        literal asterisks."""
        output_path = tmp_path / "summary.docx"
        summary_results = [
            _content_page(
                5,
                ["A point."],
                references=[
                    {
                        "citation": "North, D. C. (1997). Cliometrics. "
                        "*The American Economic Review, 87*(2), 412-414.",
                        "is_partial": False,
                    }
                ],
            )
        ]
        manager, data = build_render_context(summary_results, openalex_enabled=False)

        create_docx_summary(summary_results, output_path, "Doc", manager, data)

        doc = Document(str(output_path))
        ref_paragraph = next(
            p for p in doc.paragraphs if "American Economic Review" in p.text
        )
        assert "*" not in ref_paragraph.text
        italic_text = "".join(run.text for run in ref_paragraph.runs if run.italic)
        assert "The American Economic Review, 87" in italic_text

    def test_linked_reference_drops_emphasis_and_keeps_metadata(
        self, tmp_path: Path
    ) -> None:
        """A citation with a URL is one hyperlink without emphasis markers,
        followed by its pages and OpenAlex metadata."""
        url = "https://doi.org/10.1234/linked"
        results = [
            _content_page(
                1,
                ["A point."],
                references=[
                    {
                        "citation": "Smith, J. (2020). *A Linked Work*. Press.",
                        "is_partial": False,
                    }
                ],
            )
        ]
        manager, data = build_render_context(results, openalex_enabled=False)
        (citation,) = manager.citations.values()
        citation.url = url
        citation.doi = "10.1234/linked"
        citation.metadata = {"publication_year": 2020}
        output_path = tmp_path / "summary.docx"

        create_docx_summary(results, output_path, "Doc", manager, data)

        doc = Document(str(output_path))
        ref_paragraph = next(
            p for p in doc.paragraphs if p._p.find(_HYPERLINK) is not None
        )
        link = ref_paragraph._p.find(_HYPERLINK)
        link_text = "".join(t.text or "" for t in link.iter(_TEXT))
        assert link_text == "Smith, J. (2020). A Linked Work. Press."
        assert doc.part.rels[link.get(qn("r:id"))].target_ref == url
        runs = "".join(run.text for run in ref_paragraph.runs)
        assert runs == "[1]  (p. 1) [DOI: 10.1234/linked, Year: 2020]"


class TestDocxWriterRobustness:
    """Malformed results render, and a save never leaves a temp file behind."""

    def test_non_string_bullets_are_skipped(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.docx"
        bullets: Any = [None, 123, "real bullet"]

        create_docx_summary([_content_page(1, bullets)], output_path, "Doc")

        texts = [p.text for p in Document(str(output_path)).paragraphs]
        assert any("real bullet" in t for t in texts)
        assert not any("123" in t for t in texts)

    def test_control_character_in_heading_is_removed(self, tmp_path: Path) -> None:
        """An XML-illegal character in a page number does not fail the save."""
        output_path = tmp_path / "summary.docx"
        result = _content_page(7, ["A bullet"])
        result["page_information"]["page_number_integer"] = "7\x0b"

        create_docx_summary([result], output_path, "Doc")

        texts = [p.text for p in Document(str(output_path)).paragraphs]
        assert "Page 7" in texts
        assert not any("\x0b" in t for t in texts)

    def test_successful_save_leaves_no_temp_file(self, tmp_path: Path) -> None:
        output_path = tmp_path / "summary.docx"

        create_docx_summary([_content_page(1, ["A bullet"])], output_path, "Doc")

        assert output_path.stat().st_size > 0
        assert not output_path.with_name(output_path.name + ".tmp").exists()

    def test_failed_save_removes_temp_file(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A non-OSError during the final replace still removes the temp file."""

        def _boom(_src: Any, _dst: Any) -> None:
            raise ValueError("not an OSError")

        output_path = tmp_path / "summary.docx"
        monkeypatch.setattr("autoexcerpter.rendering.docx.os.replace", _boom)

        with pytest.raises(ValueError, match="not an OSError"):
            create_docx_summary([_content_page(1, ["A bullet"])], output_path, "Doc")

        assert [p.name for p in tmp_path.iterdir() if p.suffix == ".tmp"] == []


# ============================================================================
# Paragraph helpers
# ============================================================================
class TestAddHyperlink:
    """Tests for add_hyperlink."""

    def test_adds_hyperlink_to_paragraph(self) -> None:
        """The link run carries the text and an external relationship to the URL."""
        doc = Document()
        paragraph = doc.add_paragraph()

        add_hyperlink(paragraph, "https://example.com", "Example Link")

        link = paragraph._p.find(_HYPERLINK)
        assert link is not None
        assert "".join(t.text or "" for t in link.iter(_TEXT)) == "Example Link"
        assert doc.part.rels[link.get(qn("r:id"))].target_ref == ("https://example.com")


class TestAddFormattedTextToParagraph:
    """Tests for add_formatted_text_to_paragraph."""

    def test_plain_text_added(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()

        add_formatted_text_to_paragraph(paragraph, "Hello world")

        assert paragraph.text == "Hello world"

    def test_text_with_inline_latex(self) -> None:
        """Inline LaTeX becomes a Word equation beside the text run."""
        doc = Document()
        paragraph = doc.add_paragraph()

        add_formatted_text_to_paragraph(paragraph, "Value is $x + y$")

        assert paragraph.text == "Value is "
        assert "oMath" in paragraph._p.xml

    def test_bold_becomes_a_bold_run_without_asterisks(self) -> None:
        paragraph = Document().add_paragraph()

        add_formatted_text_to_paragraph(paragraph, "This is **bold** text")

        assert "*" not in paragraph.text
        bold_runs = [r for r in paragraph.runs if r.bold and r.text == "bold"]
        assert len(bold_runs) == 1

    def test_triple_emphasis_is_one_bold_italic_run(self) -> None:
        paragraph = Document().add_paragraph()

        add_formatted_text_to_paragraph(paragraph, "***Title***")

        assert len(paragraph.runs) == 1
        assert paragraph.runs[0].text == "Title"
        assert paragraph.runs[0].bold is True
        assert paragraph.runs[0].italic is True


# ============================================================================
# Markdown emphasis
# ============================================================================
class TestMarkdownEmphasisHelpers:
    """parse_markdown_emphasis and strip_markdown_emphasis edge cases."""

    def test_parse_mixed_emphasis(self) -> None:
        segments = parse_markdown_emphasis("a *b* c **d** e")
        assert segments == [
            ("text", "a "),
            ("italic", "b"),
            ("text", " c "),
            ("bold", "d"),
            ("text", " e"),
        ]

    def test_parse_plain_text_passthrough(self) -> None:
        assert parse_markdown_emphasis("no emphasis") == [("text", "no emphasis")]

    def test_parse_empty_string(self) -> None:
        assert parse_markdown_emphasis("") == [("text", "")]

    def test_unbalanced_asterisk_left_verbatim(self) -> None:
        assert parse_markdown_emphasis("87* (2)") == [("text", "87* (2)")]

    def test_strip_markdown_emphasis(self) -> None:
        assert (
            strip_markdown_emphasis("*The Review, 87*(2) and **bold**")
            == "The Review, 87(2) and bold"
        )

    def test_space_flanked_asterisks_are_not_italic(self) -> None:
        """CommonMark: "a * b * c" is literal prose, not emphasis."""
        segments = parse_markdown_emphasis("a * b * c")
        assert all(kind == "text" for kind, _ in segments)
        assert "".join(piece for _, piece in segments) == "a * b * c"

    def test_single_word_italic_still_parsed(self) -> None:
        assert parse_markdown_emphasis("*word*") == [("italic", "word")]

    def test_bold_still_parsed(self) -> None:
        assert parse_markdown_emphasis("**word**") == [("bold", "word")]

    def test_triple_emphasis_is_one_segment(self) -> None:
        assert parse_markdown_emphasis("***Title***") == [("bold_italic", "Title")]

    def test_triple_emphasis_inside_prose(self) -> None:
        segments = parse_markdown_emphasis("a ***b*** c")

        assert segments == [("text", "a "), ("bold_italic", "b"), ("text", " c")]
        assert all("*" not in piece for _kind, piece in segments)

    def test_strip_handles_triple_emphasis(self) -> None:
        assert strip_markdown_emphasis("***x*** and **b** and *i*") == "x and b and i"
