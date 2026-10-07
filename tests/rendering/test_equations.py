"""LaTeX parsing and simplification in rendering/equations.py.

Covers finding inline and display formulas in text, the guards that keep
LaTeX-like prose and currency as text, the placement of inline and display math
in a DOCX paragraph, whitespace normalization, the simplifications applied
before conversion and OMML sanitizing with its italic-text fallback. The OMML
the DOCX writer emits is pinned by the read-back snapshot in
test_readback_snapshots.py.
"""

from __future__ import annotations

from unittest.mock import patch

from docx import Document

import autoexcerpter.rendering.equations as equations_module
from autoexcerpter.rendering.docx import add_formatted_text_to_paragraph
from autoexcerpter.rendering.equations import (
    add_math_to_paragraph,
    normalize_latex_whitespace,
    parse_latex_in_text,
    sanitize_omml_xml,
    simplify_problematic_latex,
)

_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# A <m:groupChrPr> closed by its parent </m:groupChr>, as mathml2omml sometimes
# emits it; lxml rejects the markup.
_BROKEN_OMML = (
    f'<m:oMath xmlns:m="{_MATH_NS}"><m:groupChr><m:groupChrPr>'
    '<m:chr m:val="x"/><m:e><m:r><m:t>a</m:t></m:r></m:e>'
    "</m:groupChr></m:oMath>"
)
_REPAIRED_OMML = f'<m:oMath xmlns:m="{_MATH_NS}"><m:r><m:t>a</m:t></m:r></m:oMath>'

# Malformed markup without a groupChrPr construct, which the sanitizer cannot
# repair.
_UNREPAIRABLE_OMML = f'<m:oMath xmlns:m="{_MATH_NS}"><m:r><m:t>a</m:r></m:t></m:oMath>'


# ============================================================================
# Parsing
# ============================================================================
class TestParseLatexInText:
    """Tests for parse_latex_in_text."""

    def test_no_latex(self) -> None:
        """Text without LaTeX is one text segment."""
        result = parse_latex_in_text("Plain text without formulas")

        assert result == [("Plain text without formulas", "text")]

    def test_inline_latex(self) -> None:
        result = parse_latex_in_text("Here is $x + y = z$ inline")

        assert "latex_inline" in [kind for _, kind in result]

    def test_display_latex(self) -> None:
        result = parse_latex_in_text("Here is $$x + y = z$$ display")

        assert "latex_display" in [kind for _, kind in result]

    def test_mixed_latex(self) -> None:
        kinds = [kind for _, kind in parse_latex_in_text("Inline $a$ and $$b$$")]

        assert "latex_inline" in kinds
        assert "latex_display" in kinds

    def test_escaped_dollar(self) -> None:
        """Escaped dollar signs stay text."""
        result = parse_latex_in_text("Price is \\$50 and \\$100")

        assert "$" in "".join(content for content, _ in result)

    def test_empty_text(self) -> None:
        assert parse_latex_in_text("") == [("", "text")]

    def test_multiline_display_latex(self) -> None:
        result = parse_latex_in_text("$$\nx + y\n= z\n$$")

        assert "latex_display" in [kind for _, kind in result]


class TestLatexProseGuards:
    """Prose that merely resembles LaTeX must survive DOCX rendering."""

    def test_currency_range_is_not_inline_math(self) -> None:
        """A "$3 to $5" span is a price range, not an inline formula."""
        segments = parse_latex_in_text("Wages rose from $3 to $5 per day.")
        assert all(kind == "text" for _, kind in segments)
        assert "".join(content for content, _ in segments) == (
            "Wages rose from $3 to $5 per day."
        )

    def test_subscripted_variable_still_inline_math(self) -> None:
        segments = parse_latex_in_text("value $x_1$ here")
        assert [kind for _, kind in segments].count("latex_inline") == 1

    def test_digit_formula_with_math_indicator_kept(self) -> None:
        """A digit-initial capture carrying math markup stays math."""
        segments = parse_latex_in_text("see $2^n = m$ above")
        assert [kind for _, kind in segments].count("latex_inline") == 1

    def test_escaped_braces_survive_delimiter_sizing(self) -> None:
        simplified, _ = simplify_problematic_latex(r"\bigl\{ x \bigr\}")
        assert "\\{" in simplified
        assert "\\}" in simplified

    def test_escaped_ampersand_survives_alignment_cleanup(self) -> None:
        simplified, _ = simplify_problematic_latex(r"A \& B & C")
        assert "\\&" in simplified


class TestCurrencyDollars:
    """A currency dollar never consumes the delimiter of a genuine formula."""

    def test_math_after_currency_with_odd_dollar_count(self) -> None:
        segments = parse_latex_in_text("cost of $5 and weight of $x^2$")

        assert ("x^2", "latex_inline") in segments
        joined = "".join(content for content, kind in segments if kind == "text")
        assert "$5" in joined

    def test_currency_only_line_stays_literal(self) -> None:
        segments = parse_latex_in_text("Wages rose from $3 to $5 that year.")

        assert all(kind == "text" for _, kind in segments)

    def test_consecutive_currency_before_math(self) -> None:
        segments = parse_latex_in_text("price $5, variable $x$ and $y$")

        assert ("x", "latex_inline") in segments
        assert ("y", "latex_inline") in segments
        assert not any(
            kind == "latex_inline" and content.strip() in {"and", ", variable"}
            for content, kind in segments
        )

    def test_currency_span_with_escaped_dollar_stays_text(self) -> None:
        segments = parse_latex_in_text("cost $5 \\$ each and $x$ math")

        assert segments == [
            ("cost $5 $ each and ", "text"),
            ("x", "latex_inline"),
            (" math", "text"),
        ]

    def test_formula_with_escaped_dollar_stays_math(self) -> None:
        """A math indicator in the span wins over the currency guard."""
        segments = parse_latex_in_text("total $5 = \\$x^2$ end")

        assert segments == [
            ("total ", "text"),
            ("5 = $x^2", "latex_inline"),
            (" end", "text"),
        ]

    def test_two_inline_formulas_after_currency_and_display(self) -> None:
        """A currency dollar before a display block keeps later pairs aligned."""
        segments = parse_latex_in_text(
            "Cost $5 rose. $$E=mc^2$$ Later $x_1$ and $y_2$ done."
        )

        assert ("x_1", "latex_inline") in segments
        assert ("y_2", "latex_inline") in segments
        assert (" and ", "text") in segments

    def test_inline_formula_after_currency_and_display(self) -> None:
        segments = parse_latex_in_text("price $5 up. $$a=b$$ then $x^2$ end.")

        assert ("x^2", "latex_inline") in segments

    def test_display_then_inline_without_currency(self) -> None:
        segments = parse_latex_in_text("Text $$E=mc^2$$ then $x_1$ end.")

        assert ("E=mc^2", "latex_display") in segments
        assert ("x_1", "latex_inline") in segments


# ============================================================================
# Placement in a DOCX paragraph
# ============================================================================
class TestInlineAndDisplayMath:
    """Inline formulas flow with the text; display formulas get m:oMathPara."""

    def test_inline_formula_is_bare_omath_between_runs(self) -> None:
        paragraph = Document().add_paragraph()
        add_formatted_text_to_paragraph(paragraph, "The value $x^2$ is small.")

        xml = paragraph._p.xml
        assert "oMath" in xml
        assert "oMathPara" not in xml
        texts = [run.text for run in paragraph.runs]
        assert "The value " in texts
        assert " is small." in texts

    def test_display_formula_keeps_omathpara_wrapper(self) -> None:
        paragraph = Document().add_paragraph()
        add_formatted_text_to_paragraph(paragraph, "$$E = mc^2$$")

        assert "oMathPara" in paragraph._p.xml

    def test_mixed_line_wraps_only_the_display_formula(self) -> None:
        paragraph = Document().add_paragraph()
        add_formatted_text_to_paragraph(
            paragraph, "Given $x + y$ we obtain $$E = mc^2$$ at last."
        )

        # One opening and one closing tag.
        assert paragraph._p.xml.count("oMathPara") == 2


# ============================================================================
# Normalization and simplification
# ============================================================================
class TestNormalizeLatexWhitespace:
    """Tests for normalize_latex_whitespace."""

    def test_strips_whitespace(self) -> None:
        assert normalize_latex_whitespace("  x + y  ") == "x+y"

    def test_collapses_spaces(self) -> None:
        assert "   " not in normalize_latex_whitespace("x   +   y")

    def test_removes_space_around_operators(self) -> None:
        assert normalize_latex_whitespace("x = y + z") == "x=y+z"


class TestSimplifyProblematicLatex:
    """Tests for simplify_problematic_latex."""

    def test_logical_operators_replaced(self) -> None:
        simplified, _ = simplify_problematic_latex("a \\land b \\lor c")

        assert "\\land" not in simplified
        assert "\\lor" not in simplified

    def test_text_commands_replaced(self) -> None:
        """\\text becomes \\mathrm."""
        simplified, applied = simplify_problematic_latex("\\text{hello}")

        assert simplified == "\\mathrm{hello}"
        assert applied == ["text -> mathrm"]

    def test_left_right_delimiters_simplified(self) -> None:
        simplified, _ = simplify_problematic_latex("\\left( x \\right)")

        assert "\\left" not in simplified
        assert "\\right" not in simplified

    def test_phantom_removed(self) -> None:
        simplified, _ = simplify_problematic_latex("x + \\phantom{hidden} y")

        assert "\\phantom" not in simplified

    def test_spacing_normalized(self) -> None:
        simplified, _ = simplify_problematic_latex("x\\,y\\;z")

        assert "\\," not in simplified
        assert "\\;" not in simplified

    def test_environment_unwrapped(self) -> None:
        simplified, _ = simplify_problematic_latex("\\begin{align}x = y\\end{align}")

        assert "\\begin{align}" not in simplified
        assert "\\end{align}" not in simplified

    def test_overbrace_preserves_trailing_content(self) -> None:
        """Simplifying \\overbrace keeps the formula content after it."""
        out, _ = simplify_problematic_latex(r"\overbrace{a+b} + \frac{c}{d}")

        assert "\\frac{c}{d}" in out
        assert "a+b" in out
        assert "{d}" not in out.replace("\\frac{c}{d}", "")

    def test_underbrace_preserves_trailing_content(self) -> None:
        out, _ = simplify_problematic_latex(r"\underbrace{a+b}_{k} + \frac{c}{d}")

        assert "\\frac{c}{d}" in out
        assert "a+b" in out


# ============================================================================
# OMML sanitizing
# ============================================================================
class TestSanitizeOmmlXml:
    """Tests for sanitize_omml_xml."""

    def test_valid_xml_unchanged(self) -> None:
        valid_xml = (
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/'
            '2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>'
        )
        assert sanitize_omml_xml(valid_xml) == valid_xml

    def test_invalid_xml_returns_original(self) -> None:
        """Unfixable invalid XML returns the original string."""
        invalid_xml = "<unclosed><tags"
        assert sanitize_omml_xml(invalid_xml) == invalid_xml


class TestOmmlRecovery:
    """Markup that fails to parse is sanitized and retried before falling back."""

    def test_sanitized_markup_is_used_instead_of_italic_fallback(self) -> None:
        paragraph = Document().add_paragraph()
        seen: list[str] = []

        def fake_sanitize(markup: str) -> str:
            seen.append(markup)
            return _REPAIRED_OMML

        with (
            patch(
                "autoexcerpter.rendering.equations.mathml2omml.convert",
                return_value=_BROKEN_OMML,
            ),
            patch.object(
                equations_module, "sanitize_omml_xml", side_effect=fake_sanitize
            ),
        ):
            add_math_to_paragraph(paragraph, "x^2")

        assert len(seen) == 1
        xml = paragraph._p.xml
        assert "oMath" in xml
        assert "oMathPara" not in xml
        assert "Cambria Math" not in xml
        assert paragraph.runs == []

    def test_real_sanitizer_repairs_groupchr_markup_end_to_end(self) -> None:
        """The sanitizer closes the missing </m:groupChrPr> tag."""
        paragraph = Document().add_paragraph()

        with patch(
            "autoexcerpter.rendering.equations.mathml2omml.convert",
            return_value=_BROKEN_OMML,
        ):
            add_math_to_paragraph(paragraph, "x^2")

        xml = paragraph._p.xml
        assert "oMath" in xml
        assert "oMathPara" not in xml
        assert "Cambria Math" not in xml
        assert paragraph.runs == []

    def test_unrepairable_markup_degrades_to_italic_text(self) -> None:
        paragraph = Document().add_paragraph()

        with patch(
            "autoexcerpter.rendering.equations.mathml2omml.convert",
            return_value=_UNREPAIRABLE_OMML,
        ):
            add_math_to_paragraph(paragraph, "x^2")

        assert paragraph.runs
        assert paragraph.runs[0].italic is True
        assert paragraph.runs[0].font.name == "Cambria Math"
